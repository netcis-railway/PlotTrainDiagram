"""Read display metadata without changing event identities or saved SEQ values."""

import copy
import json
from collections import defaultdict
from zipfile import ZipFile

from openpyxl import load_workbook

from app.ingest.base_parser import integer
from app.ingest.inspection import HEADERS


def positive(value):
    try:
        return integer(value)
    except (ValueError, TypeError, OverflowError):
        return None


def lines(value):
    values = value if isinstance(value, list) else str(value or "").split(",")
    return sorted({number for v in values if (number := positive(v)) is not None})


def route_nodes(display_nodes, infrastructure):
    """Keep non-visited route positions while mapping visits to this object's own IDs."""
    routes = {}
    for item in infrastructure.get("nodes", []):
        name = item["node_name"]
        routes[name] = {
            "id": item.get("id") or f"route:{name}",
            "node_name": name,
            "node_name_chn": item.get("node_name_chn", name),
            "node_ind": positive(item.get("node_ind", item.get("NODE_IND"))),
            "line_inds": [],
            "line_seqs": {},
        }
    for member in infrastructure.get("members", []):
        name = member.get("node_name", member.get("NODE_NAME"))
        if not name:
            continue
        if name not in routes:
            routes[name] = {
                "id": f"route:{name}",
                "node_name": name,
                "node_name_chn": member.get("NODE_NAME_CHN", name),
                "node_ind": positive(member.get("node_ind", member.get("NODE_IND"))),
                "line_inds": [],
                "line_seqs": {},
            }
        item = routes[name]
        line = positive(member.get("line_ind", member.get("LINE_IND")))
        seq = positive(member.get("NODE_SEQ_OF_LINE"))
        if line is not None:
            key = str(line)
            previous = item["line_seqs"].get(key, seq)
            item["line_seqs"][key] = seq if previous == seq else None
        item["line_inds"] = sorted(
            set(item["line_inds"]) | set(lines(member.get("line_ind", member.get("LINE_IND"))))
        )
    for node in display_nodes:
        name = node["node_name"]
        # Event IDs belong to the edited object; IDs from its inherited base may differ.
        routes[name] = routes.get(name, {"line_seqs": {}}) | {
            k: copy.deepcopy(node[k])
            for k in ("id", "node_name", "node_name_chn", "node_ind", "line_inds")
        }
    return list(routes.values())


def editor_order(connection, base, kind, object_id, doc, infrastructure):
    nodes = copy.deepcopy(doc["nodes"])
    known = {n["node_name"]: n for n in infrastructure.get("nodes", [])}
    if kind == "run":
        for (data,) in connection.execute(
            "SELECT data FROM run_nodes WHERE run_id=?", [object_id]
        ).fetchall():
            item = json.loads(data)
            name = item.get("node_name", item.get("NODE_NAME"))
            if name:
                known[name] = known.get(name, {}) | item
    memberships = defaultdict(set)
    for member in infrastructure.get("members", []):
        name = member.get("node_name", member.get("NODE_NAME"))
        memberships[name].update(lines(member.get("line_ind", member.get("LINE_IND"))))
    for event in doc["events"]:
        memberships[event["node_name"]].update(lines(event.get("line_inds", [])))
    for node in nodes:
        item = known.get(node["node_name"], {}) | node
        node["node_ind"] = positive(item.get("node_ind", item.get("NODE_IND")))
        node["line_inds"] = sorted(memberships[node["node_name"]])
        if "source" in item:
            node["source"] = item["source"]

    table_rows, member_rows, course_order = [], [], []
    excel_metadata = {}
    if kind == "base":
        sources = connection.execute(
            "SELECT managed_name FROM base_tt_sources WHERE base_tt_id=? ORDER BY id", [object_id]
        ).fetchall()
    else:
        sources = connection.execute(
            "SELECT s.managed_name FROM source_files s LEFT JOIN import_jobs j ON j.id=s.job_id "
            "WHERE s.run_id=? ORDER BY j.created_at NULLS LAST,s.filename,s.id",
            [object_id],
        ).fetchall()
    complete, found_ttt = True, False
    for (managed,) in sources:
        path = (base.sources / managed).resolve()
        if path.parent != base.sources.resolve() or not path.is_file():
            complete = False
            continue
        workbook = None
        try:
            with ZipFile(path) as archive:
                if sum(item.file_size for item in archive.infolist()) > 128 * 1024 * 1024:
                    complete = False
                    continue
            workbook = load_workbook(path, read_only=True, data_only=False)
            for sheet in workbook:
                if (sheet.max_row or 0) * (sheet.max_column or 0) > 2_000_000:
                    complete = False
                    continue
                records = sheet.iter_rows(values_only=True)
                header = list(next(records, ()))
                if header[:5] == HEADERS:
                    found_ttt = True
                    for column in header[5::4]:
                        if isinstance(column, str) and column.endswith("_A"):
                            course = column[:-2]
                            if course not in course_order:
                                course_order.append(course)
                    for row in records:
                        if len(row) < 5 or not isinstance(row[4], str):
                            continue
                        name = row[4].strip()
                        if name not in table_rows:
                            table_rows.append(name)
                            excel_metadata[name] = {
                                "node_ind": positive(row[1]),
                                "line_inds": lines(row[3]),
                            }
                elif sheet.title == "NODE_INFO" and "NODE_NAME" in header:
                    for values in records:
                        item = dict(zip(header, values))
                        name = item.get("NODE_NAME")
                        if isinstance(name, str) and name.strip() not in member_rows:
                            member_rows.append(name.strip())
        except (OSError, ValueError, KeyError, TypeError, EOFError):
            complete = False
        except Exception:
            # A missing/corrupted source must not prevent viewing or saving database data.
            complete = False
        finally:
            if workbook is not None:
                workbook.close()
    for node in nodes:
        if node["node_name"] in excel_metadata:
            metadata = excel_metadata[node["node_name"]]
            node["node_ind"] = metadata["node_ind"]
            node["line_inds"] = sorted(set(node["line_inds"]) | set(metadata["line_inds"]))

    sheet_order = list(dict.fromkeys(e.get("source", {}).get("sheet", "") for e in doc["events"]))

    def provenance(item):
        source = item.get("source", {})
        sheet = source.get("sheet", "")
        return (
            sheet_order.index(sheet) if sheet in sheet_order else len(sheet_order),
            positive(source.get("row")) or 2147483647,
        )

    fallback_nodes = sorted(nodes, key=provenance)
    names = list(dict.fromkeys(table_rows + member_rows + [n["node_name"] for n in fallback_nodes]))
    name_rank = {name: rank for rank, name in enumerate(names)}
    node_order = [n["id"] for n in sorted(nodes, key=lambda n: name_rank[n["node_name"]])]
    fallback_trains = sorted(
        doc["trains"],
        key=lambda t: (
            provenance(t)[0],
            positive(t.get("source", {}).get("first_column")) or 2147483647,
        ),
    )
    course_order.extend(
        t["train_course_id"] for t in fallback_trains if t["train_course_id"] not in course_order
    )
    course_rank = {course: rank for rank, course in enumerate(course_order)}
    train_order = [
        t["id"] for t in sorted(doc["trains"], key=lambda t: course_rank[t["train_course_id"]])
    ]
    exact = found_ttt and complete
    return nodes, {
        "nodes": node_order,
        "trains": train_order,
        "source": "excel" if exact else "provenance",
        "notice": "" if exact else "部分源Excel未读取到，无法定位的项目按保留的来源记录顺序显示。",
    }
