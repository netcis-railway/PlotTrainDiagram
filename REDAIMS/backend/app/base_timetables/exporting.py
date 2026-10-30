"""Saved-state XLSX export and selective worksheet replacement."""

import copy
import json
import posixpath
from datetime import date, datetime, time, timedelta
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from app.base_timetables.workbooks import (
    book_bytes,
    format_timetable_times,
    materialize_content,
    source_content,
    working_copy_write,
)
from app.ingest.base_service import ImportConflict
from app.ingest.tracks import excel_track_field

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
TABLES = {
    "TTT": "时刻表",
    "LINE_INFO": "线路",
    "NODE_INFO": "车站线路成员",
    "MINIMUM_RUN_TIME": "最小运行时间",
}
REQUIRED = {
    "LINE_INFO": ["LINE_IND", "LINE_CODE", "LINE_NAME_CHN"],
    "NODE_INFO": [
        "LINE_IND",
        "NODE_SEQ_OF_LINE",
        "NODE_IND",
        "NODE_NAME",
        "NODE_NAME_CHN",
        "NODE_CODE",
        "DB_TRACKS",
        "UB_TRACKS",
    ],
    "MINIMUM_RUN_TIME": [
        "LINK_START_NODE_CODE",
        "LINK_END_NODE_CODE",
        "START_ACTIVITY",
        "END_ACTIVITY",
        "MINIMUM_RUN_TIME",
        "MINIMUM_RUN_TIME_SEC",
    ],
}
KEYS = {"LINE_INFO": "lines", "NODE_INFO": "members", "MINIMUM_RUN_TIME": "minimum_run_times"}


def tag(name):
    return f"{{{NS}}}{name}"


def sheet_parts(archive):
    relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {r.attrib["Id"]: r.attrib["Target"] for r in relations}
    sheets = ET.fromstring(archive.read("xl/workbook.xml")).find(tag("sheets"))
    result = {}
    for sheet in sheets:
        target = targets[sheet.attrib[f"{{{REL}}}id"]]
        path = posixpath.normpath(target.lstrip("/") if target.startswith("/") else "xl/" + target)
        if not path.startswith("xl/"):
            raise ValueError("工作簿包含不支持的表单路径")
        result[sheet.attrib["name"]] = path
    return result


def merge_selected(original, generated, titles, *, allow_new=False, keep_layout=False):
    """Keep unselected OOXML parts byte-for-byte, including formulas and drawings."""
    try:
        with ZipFile(BytesIO(original)) as source, ZipFile(BytesIO(generated)) as fresh:
            if sum(i.file_size for i in source.infolist()) > 128 * 1024 * 1024:
                raise ValueError("源工作簿解压后超过128 MB限制")
            old_parts, new_parts = sheet_parts(source), sheet_parts(fresh)
            if not allow_new and any(title not in old_parts for title in titles):
                raise ValueError("源文件缺少所选同名表单，请核对源文件或选择另存新文件")
            styles = ET.fromstring(source.read("xl/styles.xml"))
            new_styles = ET.fromstring(fresh.read("xl/styles.xml"))
            old_xfs, new_xfs = styles.find(tag("cellXfs")), new_styles.find(tag("cellXfs"))
            old_formats = styles.find(tag("numFmts"))
            if old_formats is None:
                old_formats = ET.Element(tag("numFmts"), {"count": "0"})
                styles.insert(0, old_formats)
            formats = new_styles.find(tag("numFmts"))
            format_ids = (
                {int(f.attrib["numFmtId"]): f.attrib["formatCode"] for f in formats}
                if formats is not None
                else {}
            )
            known = {f.attrib["formatCode"]: int(f.attrib["numFmtId"]) for f in old_formats}
            next_format = max([163, *known.values()]) + 1
            resources = {}
            for attribute, collection in (
                ("fontId", "fonts"),
                ("fillId", "fills"),
                ("borderId", "borders"),
            ):
                old_items, new_items = (
                    styles.find(tag(collection)),
                    new_styles.find(tag(collection)),
                )
                known_items = {ET.tostring(item): index for index, item in enumerate(old_items)}
                resources[attribute] = {}
                for index, item in enumerate(new_items):
                    key = ET.tostring(item)
                    if key not in known_items:
                        known_items[key] = len(old_items)
                        old_items.append(copy.deepcopy(item))
                    resources[attribute][index] = known_items[key]
                old_items.set("count", str(len(old_items)))
            mapping = {}
            known_xfs = {ET.tostring(item): index for index, item in enumerate(old_xfs)}
            for index, xf in enumerate(new_xfs):
                number = int(xf.attrib.get("numFmtId", "0"))
                item = copy.deepcopy(xf)
                if number in format_ids:
                    code = format_ids[number]
                    if code not in known:
                        ET.SubElement(
                            old_formats,
                            tag("numFmt"),
                            {"numFmtId": str(next_format), "formatCode": code},
                        )
                        known[code] = next_format
                        next_format += 1
                    item.set("numFmtId", str(known[code]))
                for attribute, ids in resources.items():
                    item.set(attribute, str(ids[int(item.attrib.get(attribute, "0"))]))
                item.set("xfId", "0")
                key = ET.tostring(item)
                if key not in known_xfs:
                    known_xfs[key] = len(old_xfs)
                    old_xfs.append(item)
                mapping[index] = known_xfs[key]
            old_formats.set("count", str(len(old_formats)))
            old_xfs.set("count", str(len(old_xfs)))
            updates = {"xl/styles.xml": ET.tostring(styles, encoding="utf-8", xml_declaration=True)}
            workbook = ET.fromstring(source.read("xl/workbook.xml"))
            relations = ET.fromstring(source.read("xl/_rels/workbook.xml.rels"))
            content_types = ET.fromstring(source.read("[Content_Types].xml"))
            additions = {}
            for title in titles:
                root = ET.fromstring(fresh.read(new_parts[title]))
                for cell in root.iter(tag("c")):
                    cell.set("s", str(mapping[int(cell.attrib.get("s", "0"))]))
                if title not in old_parts:
                    sheets = workbook.find(tag("sheets"))
                    sheet_id = max(int(s.attrib["sheetId"]) for s in sheets) + 1
                    part = f"xl/worksheets/redaims{sheet_id}.xml"
                    while part in source.namelist() or part in additions:
                        sheet_id += 1
                        part = f"xl/worksheets/redaims{sheet_id}.xml"
                    rid = f"redaims{sheet_id}"
                    while any(r.attrib["Id"] == rid for r in relations):
                        rid += "x"
                    ET.SubElement(
                        sheets,
                        tag("sheet"),
                        {"name": title, "sheetId": str(sheet_id), f"{{{REL}}}id": rid},
                    )
                    ET.SubElement(
                        relations,
                        "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                        {"Id": rid, "Type": REL + "/worksheet", "Target": "/" + part},
                    )
                    ET.SubElement(
                        content_types,
                        "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                        {
                            "PartName": "/" + part,
                            "ContentType": (
                                "application/vnd.openxmlformats-officedocument."
                                "spreadsheetml.worksheet+xml"
                            ),
                        },
                    )
                    old_parts[title] = part
                    additions[part] = b""
                elif keep_layout:
                    existing = ET.fromstring(source.read(old_parts[title]))
                    for name in ("sheetData", "dimension"):
                        replacement = root.find(tag(name))
                        previous = existing.find(tag(name))
                        if replacement is not None and previous is not None:
                            index = list(existing).index(previous)
                            existing.remove(previous)
                            existing.insert(index, replacement)
                    root = existing
                updates[old_parts[title]] = ET.tostring(
                    root, encoding="utf-8", xml_declaration=True
                )
            if additions:
                updates.update(
                    {
                        "xl/workbook.xml": ET.tostring(
                            workbook, encoding="utf-8", xml_declaration=True
                        ),
                        "xl/_rels/workbook.xml.rels": ET.tostring(
                            relations, encoding="utf-8", xml_declaration=True
                        ),
                        "[Content_Types].xml": ET.tostring(
                            content_types, encoding="utf-8", xml_declaration=True
                        ),
                    }
                )
            output = BytesIO()
            with ZipFile(output, "w") as result:
                for item in source.infolist():
                    result.writestr(item, updates.get(item.filename, source.read(item.filename)))
                for name in additions:
                    result.writestr(name, updates[name])
            return output.getvalue()
    except (BadZipFile, KeyError, ET.ParseError, AttributeError) as error:
        raise ValueError("源文件不是支持的标准XLSX工作簿，请核对后重试") from error


class ExportService:
    def __init__(self, editor):
        self.editor = editor
        self.base = editor.base

    def _dataset(self, c, kind, object_id):
        meta, doc, infrastructure, revision = self.editor._load(c, kind, object_id)
        if not meta["date_base"] or not doc["events"]:
            raise ValueError("当前对象尚无可导出的正式时刻数据和基准日")
        known = {n["node_name"]: n for n in infrastructure.get("nodes", [])}
        if kind == "run":
            for (value,) in c.execute(
                "SELECT data FROM run_nodes WHERE run_id=?", [object_id]
            ).fetchall():
                node = json.loads(value)
                if node.get("node_name"):
                    known[node["node_name"]] = known.get(node["node_name"], {}) | node
        for node in doc["nodes"]:
            node.update(known.get(node["node_name"], {}) | node)
            identity = c.execute("SELECT node_code FROM nodes WHERE id=?", [node["id"]]).fetchone()
            if identity:
                node.setdefault("node_code", identity[0])
        sources = sorted(
            {
                e.get("source", {}).get("sheet")
                for e in doc["events"]
                if e.get("source", {}).get("sheet")
            }
        )
        title = sources[0] if len(sources) == 1 else "TTT"
        if title in KEYS or len(title) > 31 or any(c in title for c in "[]:*?/\\"):
            title = "TTT"
        return meta, doc, infrastructure, revision, title

    def options(self, kind, object_id):
        with self.base.database.transaction() as c:
            meta, doc, infra, revision, title = self._dataset(c, kind, object_id)
            available = [
                {
                    "key": "TTT",
                    "name": title,
                    "label": "时刻表",
                    "rows": len({e["node_id"] for e in doc["events"]}),
                }
            ]
            for key, name in KEYS.items():
                if name in infra and (kind == "base" or meta["effective_base"]["base_tt_id"]):
                    available.append(
                        {"key": key, "name": key, "label": TABLES[key], "rows": len(infra[name])}
                    )
            return {
                "revision": revision,
                "sheets": available,
                "defaults": [s["key"] for s in available],
                "filename": self.filename(meta),
            }

    @staticmethod
    def filename(meta):
        name = "".join("_" if c in '<>:"/\\|?*' or ord(c) < 32 else c for c in meta["name"])
        return f"{name[:100] or '时刻表'}_{meta['date_base']}.xlsx"

    def render(self, kind, object_id, revision, selected, template=None):
        if (
            not selected
            or len(selected) != len(set(selected))
            or any(s not in TABLES for s in selected)
        ):
            raise ValueError("请选择至少一项有效且不重复的导出表单")
        epoch = None
        if template:
            if len(template) > 32 * 1024 * 1024:
                raise ValueError("源工作簿超过32 MB限制")
            try:
                with ZipFile(BytesIO(template)) as archive:
                    if sum(i.file_size for i in archive.infolist()) > 128 * 1024 * 1024:
                        raise ValueError("源工作簿解压后超过128 MB限制")
                book = load_workbook(BytesIO(template), read_only=True, data_only=False)
                epoch = book.epoch
                book.close()
            except (BadZipFile, KeyError, ET.ParseError, InvalidFileException) as error:
                raise ValueError("请选择有效的XLSX源文件") from error
        with working_copy_write(self.base) as writer, self.base.database.transaction() as c:
            meta, doc, infra, current, title = self._dataset(c, kind, object_id)
            if current != revision:
                raise ImportConflict("导出期间数据、基准日或归属已变化，请重新打开导出窗口")
            content = self._snapshot(c, kind, object_id, meta, doc, infra, selected, title)
            external = None
            if template:
                generated = self._generate(meta, doc, infra, selected, title, epoch)
                external = merge_selected(
                    template, generated, [title if s == "TTT" else s for s in selected]
                )
            saved = writer.save(c, object_id, content, kind=kind)
        return external if external is not None else saved, self.filename(meta)

    def _snapshot(self, c, kind, object_id, meta, doc, infra, selected, title):
        content = source_content(self.base, c, object_id, kind=kind, required=kind == "base")
        if content is None:
            return self._generate(meta, doc, infra, selected, title)
        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 128 * 1024 * 1024:
                raise ValueError("受管Excel解压后超过128 MB限制")
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
        try:
            epoch = book.epoch
            column_order = {}
            for key in selected:
                if key in REQUIRED and key in book.sheetnames:
                    headers = [cell.value for cell in book[key][1]]
                    if set(REQUIRED[key]).issubset(headers):
                        column_order[key] = headers
        finally:
            book.close()
        if "TTT" in selected:
            has_source = (
                kind == "base"
                or c.execute("SELECT 1 FROM source_files WHERE run_id=?", [object_id]).fetchone()
            )
            if has_source:
                book = materialize_content(
                    content, doc, meta["date_base"], require_infrastructure=kind == "base"
                )
                try:
                    updated = book_bytes(book)
                finally:
                    book.close()
                titles = sorted({e["source"]["sheet"] for e in doc["events"]})
            else:
                updated = self._generate(meta, doc, infra, ["TTT"], title, epoch)
                titles = [title]
            content = merge_selected(content, updated, titles, keep_layout=True)
        foundations = [s for s in selected if s != "TTT"]
        if foundations:
            updated = self._generate(meta, doc, infra, foundations, title, epoch, column_order)
            content = merge_selected(content, updated, foundations, allow_new=True)
        return content

    def _generate(self, meta, doc, infra, selected, title, epoch=None, column_order=None):
        book = Workbook()
        book.remove(book.active)
        if epoch:
            book.epoch = epoch
        try:
            for key in selected:
                sheet = book.create_sheet(title if key == "TTT" else key)
                if key == "TTT":
                    self._ttt(sheet, meta, doc)
                    continue
                if KEYS[key] not in infra or (
                    meta["object_type"] == "run" and not meta["effective_base"]["base_tt_id"]
                ):
                    raise ValueError("当前run没有可用基本图，无法更新所选基础表")
                rows = infra[KEYS[key]]
                if key == "NODE_INFO":
                    internal = {
                        "TRACKS_SCHEMA_VERSION",
                        "DB_STOP_TRACKS",
                        "DB_MAIN_TRACK",
                        "UB_STOP_TRACKS",
                        "UB_MAIN_TRACK",
                        "UP_TRACKS",
                    }
                    rows = [
                        {k: v for k, v in row.items() if k not in internal}
                        | {
                            "DB_TRACKS": excel_track_field(row, "DB"),
                            "UB_TRACKS": excel_track_field(row, "UB"),
                        }
                        for row in rows
                    ]
                columns = list((column_order or {}).get(key, REQUIRED[key]))
                for row in rows:
                    for name in row:
                        if name.isupper() and not name.startswith("_") and name not in columns:
                            columns.append(name)
                self._write(sheet, columns, [[r.get(k) for k in columns] for r in rows])
                sheet.freeze_panes = "A2"
            return book_bytes(book)
        finally:
            book.close()

    @staticmethod
    def _write(sheet, headers, rows):
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="215878")
        sheet.auto_filter.ref = sheet.dimensions
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 22 if column > 5 else 18

    def _ttt(self, sheet, meta, doc):
        trains = doc["trains"]
        if len({t["train_course_id"] for t in trains}) != len(trains):
            raise ValueError("同车次存在多个日期实例，请先拆分后导出，不能合并覆盖")
        if len(trains) * 4 + 5 > 16384:
            raise ValueError("列车数超过Excel列数限制，请拆分后导出")
        if len({e["node_id"] for e in doc["events"]}) * (len(trains) * 4 + 5) > 2_000_000:
            raise ValueError("导出宽表超过200万单元格限制，请拆分后导出")
        headers = ["NODE_NAME_CHN", "NODE_IND", "NODE_CODE", "LINE_IND", "NODE_NAME"]
        for train in trains:
            headers.extend(
                f"{train['train_course_id']}_{suffix}" for suffix in ("A", "D", "G", "SEQ")
            )
        events = {(e["node_id"], e["train_id"]): e for e in doc["events"]}
        start = datetime.combine(date.fromisoformat(meta["date_base"]), time())
        rows = []
        for node in sorted(
            doc["nodes"], key=lambda n: (n.get("node_ind", 2147483647), n["node_name"])
        ):
            visits = [e for e in doc["events"] if e["node_id"] == node["id"]]
            lines = sorted({i for e in visits for i in e["line_inds"]})
            if not lines:
                continue
            if len({tuple(sorted(e["line_inds"])) for e in visits}) > 1:
                raise ValueError("同站不同列车线路集合不一致，暂不能无损导出宽表")
            if node.get("node_ind") is None or not node.get("node_code"):
                raise ValueError(
                    f"车站 {node['node_name_chn']} 缺少源NODE_IND或NODE_CODE，不能推测导出"
                )
            row = [
                node["node_name_chn"],
                node["node_ind"],
                node["node_code"],
                ",".join(map(str, lines)),
                node["node_name"],
            ]
            for train in trains:
                event = events.get((node["id"], train["id"]))
                if not event:
                    row.extend([None] * 4)
                    continue
                row.extend(
                    start + timedelta(seconds=event[field]["seconds"]) if event[field] else None
                    for field in ("arrival", "departure")
                )
                row.extend([event["track_code"], event["visit_seq"]])
            rows.append(row)
        self._write(sheet, headers, rows)
        format_timetable_times(sheet)
        sheet.freeze_panes = "F2"
