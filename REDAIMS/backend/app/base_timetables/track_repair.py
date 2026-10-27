"""Explicit draft-only track repair shared by imports and the timetable editor."""

import re
from collections import defaultdict
from datetime import date
from io import BytesIO
from typing import Literal

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field

from app.base_timetables.editing import EditRequest
from app.base_timetables.service import token
from app.ingest.base_parser import integer, parse_base
from app.ingest.base_service import ImportConflict
from app.ingest.inspection import inspect_workbook, train_columns
from app.ingest.repair_source import import_source
from app.ingest.tracks import read_direction


class LineChoice(BaseModel):
    model_config = ConfigDict(extra="forbid")
    line_ind: int = Field(ge=1, le=2147483647)
    directions: list[Literal["UB", "DB"]] = Field(min_length=1, max_length=1)


class TrainChoice(BaseModel):
    model_config = ConfigDict(extra="forbid")
    train_id: str
    lines: list[LineChoice] = Field(min_length=1, max_length=1000)


class RepairRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    repair_token: str
    repair_conflicts: bool = False
    selections: list[TrainChoice] = Field(min_length=1, max_length=10000)


class EditorRepairRequest(EditRequest):
    repair_token: str
    repair_conflicts: bool = False
    selections: list[TrainChoice] = Field(min_length=1, max_length=10000)


def default_direction(course):
    first = course.split("/", 1)[0]
    match = re.search(r"(\d+)$", first)
    return ("DB" if int(match[1]) % 2 else "UB") if match else None


def repair_context(doc, infrastructure, missing=()):
    missing = set(missing)
    lines = sorted(infrastructure.get("lines", []), key=lambda line: int(line["LINE_IND"]))
    members = infrastructure.get("members", [])
    counts, warnings = conflict_context(doc)
    trains = []
    for train in doc["trains"]:
        visits = [e for e in doc["events"] if e["train_id"] == train["id"]]
        stations = {e["node_name"] for e in visits}
        related = sorted({int(m["line_ind"]) for m in members if m["node_name"] in stations})
        empty = sum(missing_track(e) for e in visits)
        trains.append(
            {
                "id": train["id"],
                "train_course_id": train["train_course_id"],
                "related_lines": related,
                "direction": default_direction(train["train_course_id"]),
                "missing_g": train["id"] in missing,
                "empty_count": empty,
                "conflict_count": counts[train["id"]],
                "visits": len(visits),
            }
        )
    return {
        "warnings": warnings,
        "trains": trains,
        "lines": [
            {"line_ind": int(line["LINE_IND"]), "line_name_chn": line["LINE_NAME_CHN"]}
            for line in lines
        ],
    }


def interval(event):
    a, d = event.get("arrival"), event.get("departure")
    if not a and not d:
        raise ValueError("缺少到达和出发时刻，无法确定占用时间")
    start = (a or d)["seconds"]
    end = (d or a)["seconds"]
    if end < start:
        raise ValueError("出发早于到达，无法分配股道，请先修正时刻")
    # A point visit reserves one second, matching the timetable's second precision.
    return start, max(start + 1, end)


def overlaps(left, right):
    return left[0] < right[1] and right[0] < left[1]


def missing_track(event):
    value = event.get("track_code")
    return value is None or (isinstance(value, str) and not value.strip())


def ordered_events(doc):
    # The same train order is shown in the repair context/window.
    order = {train["id"]: index for index, train in enumerate(doc["trains"])}
    return sorted(
        doc["events"],
        key=lambda event: (order[event["train_id"]], event.get("visit_seq", 0)),
    )


def repair_problem(doc, event, message):
    names = {
        node["node_name"]: node.get("node_name_chn", node["node_name"]) for node in doc["nodes"]
    }
    return {
        "event_id": event["id"],
        "train_course_id": event["train_course_id"],
        "node_name": event["node_name"],
        "node_name_chn": names.get(event["node_name"], event["node_name"]),
        "message": message,
    }


def existing_occupations(doc, stations=None):
    records, warnings = [], []
    for event in ordered_events(doc):
        if (stations is not None and event["node_name"] not in stations) or missing_track(event):
            continue
        try:
            track, span = integer(event["track_code"]), interval(event)
        except (ValueError, TypeError) as error:
            warnings.append(repair_problem(doc, event, f"保留已有股道，无法核查占用：{error}"))
            continue
        records.append({"event": event, "track": track, "span": span})
    return records, warnings


def conflict_message(record, other):
    return (
        f"{record['event']['train_course_id']}在本站占用股道{record['track']}，"
        f"与{other['event']['train_course_id']}的股道占用冲突"
    )


def conflict_context(doc):
    records, warnings = existing_occupations(doc)
    seen, counts = defaultdict(list), defaultdict(int)
    for record in records:
        event = record["event"]
        key = (event["node_name"], record["track"])
        for other in seen[key]:
            if overlaps(record["span"], other["span"]):
                warnings.append(repair_problem(doc, event, conflict_message(record, other)))
                counts[event["train_id"]] += 1
                counts[other["event"]["train_id"]] += 1
        seen[key].append(record)
    return counts, warnings


def repair_candidates(doc, infrastructure, event, selected, warnings):
    preferred, fallback, unknown = set(), set(), False
    passing = bool(
        event.get("arrival")
        and event.get("departure")
        and event["arrival"]["seconds"] == event["departure"]["seconds"]
    )
    for choice in selected.lines:
        for member in infrastructure.get("members", []):
            if (
                member["node_name"] != event["node_name"]
                or int(member["line_ind"]) != choice.line_ind
            ):
                continue
            try:
                tracks = read_direction(member, choice.directions[0])
            except ValueError:
                unknown = True
                continue
            stops = set(tracks.stop_tracks)
            main = {tracks.main_track} if tracks.main_track is not None else set()
            preferred.update(main if passing else stops)
            fallback.update(stops if passing else main)
    candidates = sorted(preferred) + sorted(fallback - preferred)
    if unknown:
        warnings.append(
            repair_problem(
                doc, event, "所选线路/方向的股道资料未知；仅使用已知资料，无法修复时保留原值"
            )
        )
    elif not candidates:
        warnings.append(repair_problem(doc, event, "所选线路/方向在本站没有可用股道，保留原值"))
    return candidates


def allocate(doc, infrastructure, selections, repair_conflicts=False):
    """Fill missing tracks; optionally move later conflicting existing visits in window order."""
    trains = {t["id"]: t for t in doc["trains"]}
    choices = {}
    known_lines = {int(line["LINE_IND"]) for line in infrastructure.get("lines", [])}
    for choice in selections:
        if choice.train_id not in trains or choice.train_id in choices:
            raise ValueError("修复列车不存在或重复，请重新选择")
        ids = [line.line_ind for line in choice.lines]
        if len(ids) != len(set(ids)) or not set(ids) <= known_lines:
            raise ValueError("所选线路不存在或重复")
        choices[choice.train_id] = choice
    stations = {e["node_name"] for e in doc["events"] if e["train_id"] in choices}
    records, warnings = existing_occupations(doc, stations)
    issues, assignments = [], {}
    kept = defaultdict(list)
    # Protect original occupations, including later and unselected visits.
    reserved = defaultdict(list)
    for record in records:
        reserved[(record["event"]["node_name"], record["track"])].append(record)
    for record in records:
        event = record["event"]
        selected = choices.get(event["train_id"])
        candidates = (
            repair_candidates(doc, infrastructure, event, selected, warnings) if selected else []
        )
        key = (event["node_name"], record["track"])
        conflicts = [r for r in kept[key] if overlaps(record["span"], r["span"])]
        messages = [conflict_message(record, other) for other in conflicts]
        if conflicts and repair_conflicts and selected:
            free = [
                track
                for track in candidates
                if not any(
                    r is not record and overlaps(record["span"], r["span"])
                    for r in reserved[(event["node_name"], track)]
                )
            ]
            if free:
                reserved[key].remove(record)
                record["track"] = free[0]
                assignments[event["id"]] = free[0]
                reserved[(event["node_name"], free[0])].append(record)
                messages = [message + f"；已改分配到股道{free[0]}" for message in messages]
            elif candidates:
                warnings.append(
                    repair_problem(doc, event, "股道容量不足，无法修复已有冲突，保留原股道")
                )
        for message in messages:
            warnings.append(repair_problem(doc, event, message))
        kept[(event["node_name"], record["track"])].append(record)
    jobs = defaultdict(list)
    for event in ordered_events(doc):
        if event["node_name"] not in stations or not missing_track(event):
            continue
        selected = choices.get(event["train_id"])
        if not selected:
            warnings.append(repair_problem(doc, event, "未选列车股道缺失，无法核查其占用"))
            continue
        try:
            span = interval(event)
        except ValueError as error:
            issues.append(repair_problem(doc, event, str(error)))
            continue
        candidates = repair_candidates(doc, infrastructure, event, selected, warnings)
        if candidates:
            allowed = [
                track
                for track in candidates
                if not any(overlaps(span, r["span"]) for r in reserved[(event["node_name"], track)])
            ]
            jobs[event["node_name"]].append((event, span, allowed))
    for rows in jobs.values():
        accepted, station_assignments, attempts = [], {}, 0

        def solve(remaining, occupied, trial):
            nonlocal attempts
            attempts += 1
            if attempts > 20000:
                return False
            if not remaining:
                return True
            options = [
                (
                    event,
                    span,
                    [t for t in tracks if not any(overlaps(span, s) for s in occupied[t])],
                )
                for event, span, tracks in remaining
            ]
            event, span, free = min(options, key=lambda row: len(row[2]))
            rest = [row for row in remaining if row[0]["id"] != event["id"]]
            for track in free:
                occupied[track].append(span)
                trial[event["id"]] = track
                if solve(rest, occupied, trial):
                    return True
                occupied[track].pop()
                trial.pop(event["id"], None)
            return False

        for row in rows:
            event = row[0]
            if len(accepted) >= 500 or attempts > 20000:
                message = "分配达到搜索/访问上限，保留原值，请分批修复"
            else:
                occupied = defaultdict(list)
                for earlier, span, _ in accepted:
                    occupied[station_assignments[earlier["id"]]].append(span)
                free = [
                    track
                    for track in row[2]
                    if not any(overlaps(row[1], span) for span in occupied[track])
                ]
                trial = dict(station_assignments) if free else {}
                if free:
                    trial[event["id"]] = free[0]
                if free or (row[2] and solve(accepted + [row], defaultdict(list), trial)):
                    accepted.append(row)
                    station_assignments = trial
                    continue
                message = "股道容量不足，当前没有可用且不冲突的股道，保留原值"
            warnings.append(repair_problem(doc, event, message))
        assignments.update(station_assignments)
    if issues:
        return {"can_apply": False, "changes": [], "issues": issues, "warnings": warnings}
    changes = [
        {"event_id": event["id"], "field": "track_code", "value": assignments[event["id"]]}
        for event in doc["events"]
        if event["id"] in assignments
    ]
    return {"can_apply": True, "changes": changes, "issues": [], "warnings": warnings}


class TrackRepairService:
    def __init__(self, editor):
        self.editor = editor
        self.base = editor.base

    def editor_state(self, c, kind, object_id, request):
        meta, doc, infrastructure, revision = self.editor._load(c, kind, object_id)
        if request.revision != revision:
            raise ImportConflict("数据、归属或基本图已变化，请重新打开修复窗口")
        if request.changes or (request.date_base and str(request.date_base) != meta["date_base"]):
            result, doc = self.editor._validate(c, kind, object_id, request)
            errors = [
                i
                for i in result["issues"]
                if i["severity"] == "error" and i.get("field") != "track_code"
            ]
            if errors:
                raise ValueError("请先修正其他时刻表错误：" + errors[0]["message"])
        return (
            doc,
            infrastructure,
            token([revision, request.date_base, [v.model_dump() for v in request.changes], doc]),
        )

    def editor_context(self, kind, object_id, request):
        with self.base.database.transaction() as c:
            doc, infrastructure, stamp = self.editor_state(c, kind, object_id, request)
            return repair_context(doc, infrastructure) | {"repair_token": stamp}

    def editor_apply(self, kind, object_id, request):
        with self.base.database.transaction() as c:
            doc, infrastructure, stamp = self.editor_state(c, kind, object_id, request)
            if request.repair_token != stamp:
                raise ImportConflict("修复预览已过期，请重新打开窗口")
            return allocate(doc, infrastructure, request.selections, request.repair_conflicts)

    def workbook_state(self, c, job_id):
        original, options, infrastructure, stamp = import_source(c, self.base, job_id)
        structure = inspect_workbook(original, options.get("sheet"))
        if not structure["sheets"]:
            raise ValueError("请先指定时刻表表单")
        book = load_workbook(BytesIO(original))
        missing = []
        try:
            for sheet in structure["sheets"]:
                groups, issues = train_columns(sheet["headers"])
                if any(group["missing_seq"] for group in groups):
                    raise ValueError("请先点击修复SEQ信息补齐_SEQ列，再修复股道信息")
                if sheet["headers"][:5] != [
                    "NODE_NAME_CHN",
                    "NODE_IND",
                    "NODE_CODE",
                    "LINE_IND",
                    "NODE_NAME",
                ] or any(not i.get("field", "").endswith("_G") for i in issues):
                    raise ValueError("除缺少_G以外仍有表头错误，请先修正列名与顺序")
                table = book[sheet["name"]]
                for group in reversed(groups):
                    if group["missing_g"]:
                        table.insert_cols(group["start"] + 2)
                        table.cell(1, group["start"] + 2, group["course"] + "_G")
                        missing.append(group["course"])
            normalized = BytesIO()
            book.save(normalized)
            # Clear only G values in an analysis copy, allowing invalid G values to be repaired.
            for sheet in structure["sheets"]:
                table = book[sheet["name"]]
                for col in range(8, table.max_column + 1, 4):
                    for r in range(2, table.max_row + 1):
                        table.cell(r, col).value = None
            analysis = BytesIO()
            book.save(analysis)
        finally:
            book.close()
        doc = parse_base(
            analysis.getvalue(),
            date.fromisoformat(options["date_base"]),
            options["date_mode"],
            options.get("sheet"),
            track_lists=True,
            result_context=infrastructure,
        )
        errors = [i for i in doc["issues"] if i["severity"] == "error"]
        if errors:
            raise ValueError("请先修正其他导入错误：" + errors[0]["message"])
        source = load_workbook(BytesIO(normalized.getvalue()), read_only=True)
        try:
            values = {
                s.title: list(s.values)
                for s in source
                if s.title in {v["name"] for v in structure["sheets"]}
            }
            for train in doc["trains"]:
                train["id"] = train["train_course_id"]
            for index, event in enumerate(doc["events"]):
                event["id"] = str(index)
                event["train_id"] = event["train_course_id"]
                src = event["source"]
                event["track_code"] = values[src["sheet"]][src["row"] - 1][src["first_column"] + 1]
        finally:
            source.close()
        return (
            original,
            normalized.getvalue(),
            doc,
            infrastructure if infrastructure is not None else doc,
            missing,
            stamp,
            options,
        )

    def workbook_context(self, job_id):
        with self.base.database.transaction() as c:
            _, _, doc, infrastructure, missing, stamp, _ = self.workbook_state(c, job_id)
            return repair_context(doc, infrastructure, missing) | {"repair_token": stamp}

    def workbook_apply(self, job_id, request):
        with self.base.database.transaction() as c:
            original, normalized, doc, infrastructure, _, stamp, options = self.workbook_state(
                c, job_id
            )
            if request.repair_token != stamp:
                raise ImportConflict("源文件或基本图上下文已变化，请重新打开修复窗口")
            result = allocate(doc, infrastructure, request.selections, request.repair_conflicts)
            if not result["can_apply"]:
                return result, None
            # Write selected G columns only. Nonselected missing-G groups remain blocking.
            selected = {s.train_id for s in request.selections}
            book = load_workbook(BytesIO(original))
            try:
                structure = inspect_workbook(original, options.get("sheet"))
                for sheet in structure["sheets"]:
                    table = book[sheet["name"]]
                    groups, _ = train_columns(sheet["headers"])
                    for group in reversed(groups):
                        if group["missing_g"] and group["course"] in selected:
                            table.insert_cols(group["start"] + 2)
                            table.cell(1, group["start"] + 2, group["course"] + "_G")
                assigned = {change["event_id"]: change["value"] for change in result["changes"]}
                for event in doc["events"]:
                    if event["id"] not in assigned:
                        continue
                    source = event["source"]
                    table = book[source["sheet"]]
                    headers = [cell.value for cell in table[1]]
                    col = headers.index(event["train_course_id"] + "_G") + 1
                    table.cell(source["row"], col, assigned[event["id"]])
                output = BytesIO()
                book.save(output)
                return result, output.getvalue()
            finally:
                book.close()
