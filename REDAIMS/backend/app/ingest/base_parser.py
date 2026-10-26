"""Shared V4a validation; no writes or inferred overnight rollover."""

import math
import re
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO
from xml.etree.ElementTree import ParseError

from app.ingest.inspection import HEADERS, has_time, inspect_workbook
from app.ingest.issues import node_context
from app.ingest.sequences import sequence_problem
from app.ingest.tracks import (
    TrackFormatError,
    check_event_track,
    invalid_track_message,
    normalize_track_fields,
    parse_direction,
)
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter


def raw(value):
    if isinstance(value, (datetime, date, time, timedelta)):
        return str(value)
    return value


def integer(value, minimum=1):
    if isinstance(value, bool):
        raise ValueError("须为整数")
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation as error:
        raise ValueError("须为整数") from error
    if not number.is_finite() or number != number.to_integral_value() or number < minimum:
        raise ValueError(f"须为不小于 {minimum} 的整数")
    if number > 2147483647:
        raise ValueError("整数超出支持范围")
    return int(number)


def seconds(value):
    if isinstance(value, timedelta):
        total = Decimal(str(value.total_seconds()))
    elif isinstance(value, time):
        total = Decimal(value.hour * 3600 + value.minute * 60 + value.second)
        total += Decimal(value.microsecond) / 1000000
    elif isinstance(value, str):
        match = re.fullmatch(r"(\d+):(\d{2}):(\d{2}(?:\.\d+)?)", value.strip())
        if not match:
            raise ValueError("须为 HH:MM:SS，可使用超过24小时的累计小时")
        hour, minute, second = Decimal(match[1]), Decimal(match[2]), Decimal(match[3])
        if minute >= 60 or second >= 60:
            raise ValueError("分、秒须小于60")
        total = hour * 3600 + minute * 60 + second
    else:
        raise ValueError("时间须为 Excel 日期/时间或 HH:MM:SS 文本；裸数值含义不明确")
    if total < 0:
        raise ValueError("时刻不能为负")
    return int(total.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def event_time(value, date_base, origin, mode):
    if value is None or value == "":
        return None
    offset = 0
    if isinstance(value, datetime):
        if mode == "preserve" and value.year < 2020:
            raise ValueError("源日期早于2020年，请选择重新指定日期并确认")
        offset = (value.date() - (origin if mode == "rebase" else date_base)).days
        value = value.time()
    total = offset * 86400 + seconds(value)
    if total < 0:
        raise ValueError("时刻早于运行基准日，请调整基准日")
    day, remainder = divmod(total, 86400)
    try:
        date_base + timedelta(days=day)
    except OverflowError as error:
        raise ValueError("时间超出可表示的日历日期范围") from error
    h, remainder = divmod(remainder, 3600)
    m, sec = divmod(remainder, 60)
    return {"hms": f"{h:02}:{m:02}:{sec:02}", "seconds": total, "day_offset": day}


def track_list(value):
    return parse_direction(value).numbers


def parse_base(
    content: bytes,
    date_base: date,
    mode: str,
    selected_sheet: str | None = None,
    *,
    track_lists: bool = False,
    result_context: dict | None = None,
):
    structure = inspect_workbook(content, selected_sheet)
    issues = []
    header_cache = {}
    source_nodes = {}

    def issue(sheet, row, field, value, message, severity="error", column=None):
        if column is None and field in header_cache.get(sheet, []):
            column = header_cache[sheet].index(field) + 1
        issues.append(
            {
                "sheet": sheet,
                "row": row,
                "field": field,
                "cell": f"{get_column_letter(column)}{row}" if column else str(row),
                "value": raw(value),
                "message": message,
                "severity": severity,
                **source_nodes.get((sheet, row), {}),
            }
        )

    if not structure["sheets"]:
        issue("", 1, "Sheet", None, "未找到 TTT 表，请手动指定")
    for sheet in structure["sheets"]:
        for problem in sheet["issues"]:
            cell = problem.get("cell", "")
            letters = re.match(r"^([A-Z]+)1$", cell)
            column = None
            if letters:
                column = column_index_from_string(letters[1])
            issue(
                sheet["name"],
                1,
                problem.get("field", "header"),
                None,
                problem["message"],
                column=column,
            )
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (ParseError, ValueError) as error:
        raise ValueError("工作簿内容损坏") from error
    nodes = {}
    by_ind = {}
    identity_maps = [dict(), dict(), dict()]
    lines = {}
    members = []
    minimum = []
    events = []
    sources = []
    previews = []

    def records(title, required):
        if title not in workbook.sheetnames:
            issue(title, 1, "Sheet", None, "基本图缺少必需表")
            return []
        sheet = workbook[title]
        if (sheet.max_row or 0) * (sheet.max_column or 0) > 2_000_000:
            raise ValueError(f"{title} 超过200万单元格限制")
        iterator = sheet.iter_rows(values_only=True)
        header = list(next(iterator, ()))
        if (
            track_lists
            and title == "NODE_INFO"
            and "UB_TRACKS" not in header
            and "UP_TRACKS" in header
        ):
            header[header.index("UP_TRACKS")] = "UB_TRACKS"
        header_cache[title] = header
        for field in required:
            if field not in header:
                issue(title, 1, field, None, "缺少必需表头")
        named = [name for name in header if name is not None]
        if len(named) != len(set(named)):
            issue(title, 1, "header", None, "表头重复")
        result = []
        preview_rows = []
        for number, values in enumerate(iterator, 2):
            if not any(value is not None for value in values):
                continue
            row = {str(key): raw(value) for key, value in zip(header, values) if key is not None}
            row["_row"] = number
            row["_sheet"] = title
            source_nodes[(title, number)] = node_context(row)
            result.append(row)
            if len(preview_rows) < 200:
                preview_rows.append({"row": number, "values": [raw(v) for v in values]})
            for col, value in enumerate(values, 1):
                if isinstance(value, str) and value.startswith("="):
                    issue(
                        title,
                        number,
                        str(header[col - 1]),
                        value,
                        "暂不接受公式值，请在 Excel 中转为已确认的数值",
                        column=col,
                    )
        previews.append(
            {"name": title, "headers": header, "rows": preview_rows, "total_rows": len(result)}
        )
        return result

    def required_text(row, field):
        value = row.get(field)
        if not isinstance(value, str) or not value.strip():
            raise ValueError(f"{field} 必须是非空文本")
        return value.strip()

    def register_node(row):
        identity = tuple(
            required_text(row, key) for key in ("NODE_NAME", "NODE_NAME_CHN", "NODE_CODE")
        )
        ind = integer(row.get("NODE_IND"))
        for index, name in enumerate(identity):
            if name in identity_maps[index] and identity_maps[index][name] != identity:
                raise ValueError("节点名称、中文名、编码不是一一对应")
            identity_maps[index][name] = identity
        if ind in by_ind and by_ind[ind] != identity:
            raise ValueError("同一路网 NODE_IND 对应不同节点")
        if identity[0] in nodes and nodes[identity[0]]["node_ind"] != ind:
            raise ValueError("同一节点使用了不同 NODE_IND")
        by_ind[ind] = identity
        nodes[identity[0]] = {
            "node_ind": ind,
            "node_name": identity[0],
            "node_name_chn": identity[1],
            "node_code": identity[2],
            "source": {"sheet": row["_sheet"], "row": row["_row"]},
        }
        return identity[0]

    try:
        if result_context is None:
            line_rows = records("LINE_INFO", ["LINE_IND", "LINE_CODE", "LINE_NAME_CHN"])
            codes = set()
            for row in line_rows:
                try:
                    ind = integer(row.get("LINE_IND"))
                    code = required_text(row, "LINE_CODE")
                    required_text(row, "LINE_NAME_CHN")
                    if ind in lines or code in codes:
                        raise ValueError("LINE_IND 或 LINE_CODE 重复")
                    for field in ("NUM_NODE",):
                        if row.get(field) is not None:
                            integer(row[field], 0)
                    for field in ("SPEED_KMPH", "LENGTH_KM"):
                        if row.get(field) is not None:
                            if not math.isfinite(float(row[field])) or float(row[field]) < 0:
                                raise ValueError(f"{field} 须为非负有限数")
                    lines[ind] = row
                    codes.add(code)
                except (ValueError, TypeError) as error:
                    issue("LINE_INFO", row["_row"], "LINE_IND", row.get("LINE_IND"), str(error))
            member_rows = records(
                "NODE_INFO",
                [
                    "LINE_IND",
                    "LINE_CODE",
                    "LINE_NAME_CHN",
                    "NODE_SEQ_OF_LINE",
                    *HEADERS[:3],
                    "NODE_NAME",
                    "DB_TRACKS",
                    "UB_TRACKS",
                ],
            )
            member_keys = set()
            station_order = defaultdict(list)
            attributes = {}
            for row in member_rows:
                try:
                    node_name = register_node(row)
                    line = integer(row.get("LINE_IND"))
                    seq = integer(row.get("NODE_SEQ_OF_LINE"))
                    if line not in lines or row.get("LINE_CODE") != lines[line]["LINE_CODE"]:
                        raise ValueError("线路索引/编码与 LINE_INFO 不一致")
                    if row.get("LINE_NAME_CHN") != lines[line]["LINE_NAME_CHN"]:
                        raise ValueError("线路中文名与 LINE_INFO 不一致")
                    if (line, node_name) in member_keys or seq in station_order[line]:
                        raise ValueError("线路成员或站序重复")
                    member_keys.add((line, node_name))
                    station_order[line].append(seq)
                    track_fields = normalize_track_fields(row) if track_lists else {}
                    for field in (
                        "DB_TRACKS",
                        "UB_TRACKS",
                        "MINIMUM_PRE_STATION_TURNBACK_TIME_SECONDS",
                        "MINIMUM_POST_STATION_TURNBACK_TIME_SECONDS",
                    ):
                        if row.get(field) is not None and not (
                            track_lists and field in {"DB_TRACKS", "UB_TRACKS"}
                        ):
                            integer(row[field], 0)
                    for field, limit in (("LATITUDE", 90), ("LONGITUDE", 180)):
                        if row.get(field) is not None:
                            if (
                                not math.isfinite(float(row[field]))
                                or abs(float(row[field])) > limit
                            ):
                                raise ValueError(f"{field} 超出范围")
                    for field in (
                        "NODE_CATEGORY",
                        "LATITUDE",
                        "LONGITUDE",
                        "TURNBACK_TYPE",
                        "MINIMUM_PRE_STATION_TURNBACK_TIME_SECONDS",
                        "MINIMUM_POST_STATION_TURNBACK_TIME_SECONDS",
                        "OVERTAKING_ALLOWED",
                    ):
                        value = row.get(field)
                        key = (node_name, field)
                        if value is not None:
                            if key in attributes and attributes[key] != value:
                                raise ValueError(f"同一节点 {field} 多处非空值不一致")
                            attributes[key] = value
                    if row.get("DB_TRACKS") is None or row.get("UB_TRACKS") is None:
                        issue(
                            "NODE_INFO",
                            row["_row"],
                            "DB_TRACKS/UB_TRACKS",
                            None,
                            "到发线/正线信息缺失，暂不能展开该线路的车站股道"
                            if track_lists
                            else "股道数量缺失，暂不能展开该站股道",
                            "warning",
                        )
                    members.append(
                        {**row, **track_fields, "node_name": node_name, "line_ind": line}
                    )
                except (ValueError, TypeError) as error:
                    field = error.field if isinstance(error, TrackFormatError) else "node"
                    issue(
                        "NODE_INFO",
                        row["_row"],
                        field,
                        row.get(field, row.get("NODE_NAME")),
                        str(error),
                    )
            for line, row in lines.items():
                order = sorted(station_order[line])
                if order and order != list(range(1, len(order) + 1)):
                    issue("NODE_INFO", 1, "NODE_SEQ_OF_LINE", line, "已提供的线路站序不连续")
                if row.get("NUM_NODE") is not None and integer(row["NUM_NODE"], 0) != len(order):
                    issue(
                        "LINE_INFO",
                        row["_row"],
                        "NUM_NODE",
                        row["NUM_NODE"],
                        "声明节点数与成员数不同",
                    )
                if len(order) <= 1:
                    issue(
                        "LINE_INFO",
                        row["_row"],
                        "LINE_IND",
                        line,
                        "线路基础信息不足（不足两个节点），不能据此作图",
                        "warning",
                    )
        else:
            # Result imports never read foundation sheets from the uploaded workbook.
            lines = {int(row["LINE_IND"]): row for row in result_context.get("lines", [])}
            members = result_context.get("members", [])
            member_keys = {(int(m["line_ind"]), m["node_name"]) for m in members}
        dates = []
        ttt_rows = []
        for sheet in structure["sheets"]:
            title = sheet["name"]
            table = workbook[title]
            iterator = table.iter_rows(values_only=True)
            header = list(next(iterator, ()))
            header_cache[title] = header
            if sheet["issues"]:
                continue
            sources.append({"sheet": title, "trains": sheet["trains"]})
            preview = []
            for number, values in enumerate(iterator, 2):
                if not any(value is not None for value in values):
                    continue
                if len(preview) < 200:
                    preview.append({"row": number, "values": [raw(v) for v in values]})
                row = dict(zip(HEADERS, values[:5]))
                row.update(_row=number, _sheet=title)
                source_nodes[(title, number)] = node_context(row)
                try:
                    node_name = register_node(row)
                    line_ids = [integer(v.strip()) for v in str(row["LINE_IND"]).split(",")]
                    for line in line_ids:
                        if line not in lines:
                            if result_context is None:
                                raise ValueError(f"线路 {line} 未在 LINE_INFO 定义")
                            issue(
                                title,
                                number,
                                "LINE_IND",
                                line,
                                "有效基本图未定义该线路，保留源编号，相关绘图/分析受限",
                                "warning",
                                4,
                            )
                        if (line, node_name) not in member_keys:
                            issue(
                                title,
                                number,
                                "LINE_IND",
                                line,
                                "该节点缺少 NODE_INFO 线路成员，保留关联但不能推测站序",
                                "warning",
                                4,
                            )
                    row["line_inds"] = line_ids
                except (ValueError, TypeError) as error:
                    issue(title, number, "node", row.get("NODE_NAME"), str(error))
                    continue
                for start, course in zip(range(5, len(header), 4), sheet["trains"]):
                    group = list(values[start : start + 4])
                    if len(group) != 4 or all(v is None for v in group):
                        continue
                    for col, value in enumerate(group, start + 1):
                        if isinstance(value, str) and value.startswith("="):
                            issue(
                                title,
                                number,
                                str(header[col - 1]),
                                value,
                                "公式须转为已确认值",
                                column=col,
                            )
                    if not any(has_time(group[index]) for index in (0, 1, 3)):
                        if has_time(group[2]):
                            issue(
                                title,
                                number,
                                f"{course}_G",
                                raw(group[2]),
                                "该行无到发时刻且无SEQ，仅有股道值，不作为列车访问参与导入",
                                "warning",
                                start + 3,
                            )
                        continue
                    for value in group[:2]:
                        if isinstance(value, datetime):
                            dates.append(value.date())
                    ttt_rows.append((row, node_name, course, start, group))
            previews.append(
                {
                    "name": title,
                    "headers": header,
                    "rows": preview,
                    "total_rows": sum(1 for r in preview)
                    if table.max_row is None
                    else table.max_row - 1,
                }
            )
        origin = min(dates) if dates else date_base
        if dates and min(dates).year < 2020 <= max(dates).year:
            issue("TTT", 1, "date", None, "旧日期与现代日期混合，须先修正源文件，不能整体重置")
        if dates and origin.year < 2020:
            issue(
                "TTT",
                1,
                "date",
                origin.isoformat(),
                "源日期早于2020年，必须核对日期重置方案",
                "warning",
            )
        trains = defaultdict(list)
        event_keys = set()
        sequence_records = defaultdict(list)
        invalid_sequence_courses = set()
        for row, node_name, course, start, group in ttt_rows:
            title, number = row["_sheet"], row["_row"]
            field, error_column = f"{course}_SEQ", start + 4
            try:
                seq = integer(group[3])
            except (ValueError, TypeError, OverflowError) as error:
                invalid_sequence_courses.add(course)
                issue(title, number, field, raw(group[3]), str(error), column=error_column)
                continue
            sequence_records[course].append(
                {"seq": seq, "sheet": title, "row": number, "column": start + 4}
            )
            try:
                if (course, node_name) in event_keys:
                    raise ValueError("跨 Sheet 同车同站重复，请先在源文件确认合并结果")
                event_keys.add((course, node_name))
                field, error_column = f"{course}_G", start + 3
                track = None if group[2] is None else integer(group[2])
                matching = [
                    m
                    for m in members
                    if m["node_name"] == node_name
                    and m["line_ind"] in row["line_inds"]
                    and m.get("DB_TRACKS") is not None
                    and m.get("UB_TRACKS") is not None
                ]
                if track_lists:
                    if track is not None:
                        state, allowed, station_lines = check_event_track(members, node_name, track)
                        if state == "invalid":
                            raise ValueError(invalid_track_message(track, station_lines, allowed))
                        if state == "unknown":
                            issue(
                                title,
                                number,
                                field,
                                track,
                                "本站所属线路的股道资料不完整，无法完整核对该编号，请核验基本图",
                                "warning",
                                error_column,
                            )
                else:
                    limits = [
                        integer(m["DB_TRACKS"], 0) + integer(m["UB_TRACKS"], 0) for m in matching
                    ]
                    if track is not None and limits and track > min(limits):
                        raise ValueError("股道编号超过 NODE_INFO 的股道总数")
                field, error_column = f"{course}_A", start + 1
                arrival = event_time(group[0], date_base, origin, mode)
                field, error_column = f"{course}_D", start + 2
                departure = event_time(group[1], date_base, origin, mode)
                event = {
                    "train_course_id": course,
                    "node_name": node_name,
                    "visit_seq": seq,
                    "track_code": track,
                    "arrival": arrival,
                    "departure": departure,
                    "line_inds": row["line_inds"],
                    "source": {
                        "sheet": title,
                        "row": number,
                        "first_column": start + 1,
                        "raw": [raw(v) for v in group],
                    },
                }
                events.append(event)
                trains[course].append(event)
            except (ValueError, TypeError, OverflowError) as error:
                issue(
                    title,
                    number,
                    field,
                    raw(group[error_column - start - 1]),
                    str(error),
                    column=error_column,
                )
        fragments = defaultdict(set)
        for event in events:
            fragments[event["train_course_id"]].add(event["source"]["sheet"])
        for course, sheets in fragments.items():
            if len(sheets) > 1:
                issue(
                    "TTT",
                    1,
                    course,
                    sorted(sheets),
                    "同车次跨多个TTT表，需要先确认片段/实例合并；当前请整理为单表后导入",
                )
        for course, seq_rows in sequence_records.items():
            if course in invalid_sequence_courses:
                continue
            ordered = sorted(seq_rows, key=lambda item: item["seq"])
            problem = sequence_problem([item["seq"] for item in ordered])
            if problem:
                message, index = problem
                source = ordered[index]
                issue(
                    source["sheet"],
                    source["row"],
                    f"{course}_SEQ",
                    source["seq"],
                    message,
                    column=source["column"],
                )
        candidates = []
        for course, visits in trains.items():
            visits.sort(key=lambda e: e["visit_seq"])
            first = next((v for v in visits if v["departure"] is not None), None)
            if first is None:
                issue(
                    visits[0]["source"]["sheet"],
                    visits[0]["source"]["row"],
                    course,
                    None,
                    "全部缺少发车时刻，列车基准日无法确认",
                )
            else:
                base = date_base + timedelta(days=first["departure"]["day_offset"])
                fallback = first is not visits[0]
                if fallback:
                    issue(
                        first["source"]["sheet"],
                        first["source"]["row"],
                        course,
                        None,
                        "首条记录缺发车时刻，按有发车时刻的最小SEQ判定，须人工核验",
                        "warning",
                    )
                candidates.append(
                    {
                        "train_course_id": course,
                        "train_base_date": base.isoformat(),
                        "adopted_seq": first["visit_seq"],
                        "source": first["source"],
                        "departure": first["departure"],
                        "fallback": fallback,
                    }
                )
            for index, visit in enumerate(visits):
                if (index > 0 and visit["arrival"] is None) or (
                    index < len(visits) - 1 and visit["departure"] is None
                ):
                    issue(
                        visit["source"]["sheet"],
                        visit["source"]["row"],
                        course,
                        None,
                        "到发存在缺失，保留空值，后续绘图断线且排除相关比较",
                        "warning",
                    )
        if result_context is None:
            minimum_rows = records(
                "MINIMUM_RUN_TIME",
                [
                    "LINK_START_NODE_CODE",
                    "LINK_END_NODE_CODE",
                    "START_ACTIVITY",
                    "END_ACTIVITY",
                    "MINIMUM_RUN_TIME",
                    "MINIMUM_RUN_TIME_SEC",
                ],
            )
            codes_to_names = {n["node_code"]: n["node_name"] for n in nodes.values()}
            seen = set()
            for row in minimum_rows:
                try:
                    start = codes_to_names.get(row.get("LINK_START_NODE_CODE"))
                    end = codes_to_names.get(row.get("LINK_END_NODE_CODE"))
                    if not start or not end:
                        raise ValueError("区间端点无法解析到节点")
                    a, d = row.get("START_ACTIVITY"), row.get("END_ACTIVITY")
                    if a not in ("PASS", "STOP") or d not in ("PASS", "STOP"):
                        raise ValueError("活动类型须为 PASS 或 STOP")
                    duration = seconds(row.get("MINIMUM_RUN_TIME"))
                    if duration != integer(row.get("MINIMUM_RUN_TIME_SEC"), 0):
                        raise ValueError("主运行时间与秒值不一致")
                    key = (start, end, a, d)
                    if key in seen:
                        raise ValueError("有向区间与活动组合重复")
                    seen.add(key)
                    minimum.append(
                        {**row, "start_node_name": start, "end_node_name": end, "seconds": duration}
                    )
                except (ValueError, TypeError) as error:
                    issue(
                        "MINIMUM_RUN_TIME",
                        row["_row"],
                        "MINIMUM_RUN_TIME",
                        row.get("MINIMUM_RUN_TIME"),
                        str(error),
                    )
            for line in lines:
                ordered = sorted(
                    (m for m in members if m["line_ind"] == line),
                    key=lambda m: integer(m["NODE_SEQ_OF_LINE"]),
                )
                for left, right in zip(ordered, ordered[1:]):
                    key = (left["node_name"], right["node_name"], "PASS", "PASS")
                    if key not in seen or not any(
                        m["start_node_name"] == key[0]
                        and m["end_node_name"] == key[1]
                        and m["seconds"] > 0
                        for m in minimum
                    ):
                        issue(
                            "MINIMUM_RUN_TIME",
                            1,
                            "PASS/PASS",
                            str(key[:2]),
                            "缺少正向相邻站 PASS→PASS 正主时间，该线路暂不能作图",
                            "warning",
                        )
        if not events:
            issue("TTT", 1, "events", None, "没有有效列车访问")
        return {
            "tracks_schema_version": 2 if track_lists and result_context is None else None,
            "issues": issues,
            "date_base": date_base.isoformat(),
            "date_mode": mode,
            "source_date_origin": origin.isoformat(),
            "date_system": str(workbook.epoch),
            "nodes": list(nodes.values()),
            "lines": list(lines.values()) if result_context is None else [],
            "members": members if result_context is None else [],
            "minimum_run_times": minimum,
            "events": events,
            "trains": candidates,
            "sources": sources,
            "previews": previews,
            "available_sheets": workbook.sheetnames,
        }
    finally:
        workbook.close()
