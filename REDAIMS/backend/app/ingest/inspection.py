from io import BytesIO
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

HEADERS = ["NODE_NAME_CHN", "NODE_IND", "NODE_CODE", "LINE_IND", "NODE_NAME"]


def inspect_workbook(content: bytes, selected_sheet: str | None = None):
    """Read-only structural preview, never a successful business import."""
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 128 * 1024 * 1024:
                raise ValueError("工作簿解压后超过 128 MB 限制")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, KeyError, OSError, ParseError, InvalidFileException) as error:
        raise ValueError("无法读取 XLSX 文件") from error
    try:
        if selected_sheet and selected_sheet not in workbook.sheetnames:
            raise ValueError("指定的 Sheet 不存在")
        sheets = []
        for sheet in workbook:
            selected = (
                sheet.title == selected_sheet
                if selected_sheet
                else sheet.title.upper().startswith("TTT")
            )
            if not selected:
                continue
            if (sheet.max_row or 0) * (sheet.max_column or 0) > 2_000_000:
                raise ValueError(f"{sheet.title} 超过 200 万单元格限制")
            rows = sheet.iter_rows(values_only=True)
            header = list(next(rows, ()))
            while header and header[-1] is None:
                header.pop()
            issues = []
            if header[:5] != HEADERS:
                issues.append({"cell": "A1:E1", "message": "前五列表头不符合 TTT 格式"})
            groups, problems = train_columns(header)
            issues.extend(problems)
            trains = [group["course"] for group in groups]
            if len(trains) != len(set(trains)):
                issues.append({"cell": "F1", "message": "同一 Sheet 的车次四列组重复"})
            preview = []
            for row in rows:
                if len(preview) >= 20:
                    break
                preview.append(
                    [None if value is None else str(value) for value in row[: len(header)]]
                )
            sheets.append(
                {
                    "name": sheet.title,
                    "rows": max(0, (sheet.max_row or 1) - 1),
                    "headers": header,
                    "trains": trains,
                    "issues": issues,
                    "preview": preview,
                }
            )
        return {
            "stage": "structure_only",
            "sheets": sheets,
            "available_sheets": workbook.sheetnames,
            "requires_sheet_selection": not sheets,
            "message": "仅结构预览；日期、SEQ、跨表映射等业务校验及正式导入尚未执行",
        }
    finally:
        workbook.close()


def train_columns(header):
    """Locate groups with missing G and/or SEQ without shifting subsequent trains."""
    groups, issues = [], []
    column = 5
    while column < len(header):
        name = header[column]
        if not isinstance(name, str) or not name.endswith("_A") or name == "_A":
            issues.append(
                {"cell": f"{get_column_letter(column + 1)}1", "message": "列车四列组名称或顺序错误"}
            )
            column += 1
            continue
        course = name[:-2]
        if header[column : column + 2] != [f"{course}_A", f"{course}_D"]:
            issues.append(
                {
                    "cell": f"{get_column_letter(column + 1)}1",
                    "message": f"列车 {course} 必须按 A、D、G、SEQ 顺序排列，且到达/出发列不能缺失",
                }
            )
            column += 1
            continue
        end = column + 2
        missing_g = end >= len(header) or header[end] != f"{course}_G"
        if not missing_g:
            end += 1
        missing_seq = end >= len(header) or header[end] != f"{course}_SEQ"
        if not missing_seq:
            end += 1
        groups.append(
            {
                "course": course,
                "start": column + 1,
                "width": end - column,
                "missing_g": missing_g,
                "missing_seq": missing_seq,
            }
        )
        if missing_g:
            issues.append(
                {
                    "cell": f"{get_column_letter(column + 3)}1",
                    "field": f"{course}_G",
                    "message": f"列车 {course} 缺少股道列 {course}_G（应插在 {course}_SEQ 前），"
                    "请使用修复股道信息后重新预检",
                }
            )
        if missing_seq:
            issues.append(
                {
                    "cell": f"{get_column_letter(end + 1)}1",
                    "field": f"{course}_SEQ",
                    "message": f"列车 {course} 缺少序号列 {course}_SEQ，"
                    "请使用修复SEQ信息按源表行序生成连续编号后重新预检",
                }
            )
        column = end
    if not groups:
        issues.append({"cell": "F1", "message": "列车必须使用 A、D、G、SEQ 四列组"})
    return groups, issues


def has_time(value):
    """Presence, not time validity; invalid nonblank values remain preflight errors."""
    return value is not None and (not isinstance(value, str) or bool(value.strip()))
