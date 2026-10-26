"""Add only missing SEQ columns, numbered by source rows with actual A/D values."""

from io import BytesIO

from app.base_timetables.service import token
from app.ingest.base_service import ImportConflict
from app.ingest.inspection import HEADERS, has_time, inspect_workbook, train_columns
from app.ingest.repair_source import import_source
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field


class SequenceRepairRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    repair_token: str
    train_ids: list[str] = Field(min_length=1, max_length=10000)


def sequence_rows(table, start):
    return [
        row
        for row in range(2, table.max_row + 1)
        if has_time(table.cell(row, start).value) or has_time(table.cell(row, start + 1).value)
    ]


class SequenceRepairService:
    def __init__(self, base):
        self.base = base

    def state(self, c, job_id):
        original, options, _, stamp = import_source(c, self.base, job_id)
        structure = inspect_workbook(original, options.get("sheet"))
        if not structure["sheets"]:
            raise ValueError("请先指定时刻表表单")
        trains = []
        book = load_workbook(BytesIO(original))
        try:
            for sheet in structure["sheets"]:
                groups, _ = train_columns(sheet["headers"])
                if sheet["headers"][:5] != HEADERS or any(
                    not issue.get("field", "").endswith(("_G", "_SEQ")) for issue in sheet["issues"]
                ):
                    raise ValueError("除缺少_G/_SEQ以外仍有表头错误，请先修正列名、重复车次与顺序")
                table = book[sheet["name"]]
                for group in groups:
                    rows = sequence_rows(table, group["start"])
                    trains.append(
                        {
                            "id": token([sheet["name"], group["course"]]),
                            "sheet": sheet["name"],
                            "train_course_id": group["course"],
                            "missing_seq": group["missing_seq"],
                            "missing_g": group["missing_g"],
                            "visits": len(rows),
                            "first_row": rows[0] if rows else None,
                            "last_row": rows[-1] if rows else None,
                            "start": group["start"],
                            "width": group["width"],
                        }
                    )
            return original, trains, stamp
        finally:
            book.close()

    def context(self, job_id):
        with self.base.database.transaction() as c:
            _, trains, stamp = self.state(c, job_id)
            return {"repair_token": stamp, "trains": trains}

    def apply(self, job_id, request):
        with self.base.database.transaction() as c:
            original, trains, stamp = self.state(c, job_id)
            if request.repair_token != stamp:
                raise ImportConflict("源文件或导入上下文已变化，请重新打开修复窗口")
            selected = set(request.train_ids)
            eligible = {train["id"] for train in trains if train["missing_seq"]}
            if len(selected) != len(request.train_ids) or not selected <= eligible:
                raise ValueError("仅可选择缺少_SEQ列的列车，不能重编号已有列或重复选择")
            book = load_workbook(BytesIO(original))
            try:
                # Right to left: inserting one column cannot shift an earlier group's coordinates.
                for train in sorted(trains, key=lambda train: train["start"], reverse=True):
                    if train["id"] not in selected:
                        continue
                    table = book[train["sheet"]]
                    rows = sequence_rows(table, train["start"])
                    if not rows:
                        raise ValueError(
                            f"{train['sheet']} / {train['train_course_id']} "
                            "没有到达或出发时刻，不能生成SEQ"
                        )
                    column = train["start"] + train["width"]
                    table.insert_cols(column)
                    table.cell(1, column, train["train_course_id"] + "_SEQ")
                    for seq, row in enumerate(rows, 1):
                        table.cell(row, column, seq)
                output = BytesIO()
                book.save(output)
                return output.getvalue()
            finally:
                book.close()
