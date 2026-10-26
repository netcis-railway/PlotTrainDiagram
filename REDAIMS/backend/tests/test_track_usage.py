from copy import deepcopy
from datetime import date
from io import BytesIO

import pytest
from app.core.config import Settings
from app.diagram.tracks import expand_track_usage, track_resources
from app.ingest.base_parser import parse_base
from app.ingest.tracks import (
    excel_track_field,
    member_tracks,
    normalize_track_fields,
    parse_direction,
)
from app.main import create_app
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sample_workbook import SAMPLE_WORKBOOK
from test_base_import import commit, create_base, preview
from test_base_replace import seed_result
from test_catalog_deletion import branch


def track_book():
    book = load_workbook(BytesIO(SAMPLE_WORKBOOK.read_bytes()))
    sheet = book["NODE_INFO"]
    header = [c.value for c in sheet[1]]
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, header.index("DB_TRACKS") + 1, "[15,14,13,12,11,10]/9")
        sheet.cell(row, header.index("UB_TRACKS") + 1, "[10,11,12,13,14,15]/16")
    ttt = book["TTT1base"]
    for row in range(2, ttt.max_row + 1):
        for col in range(8, ttt.max_column + 1, 4):
            if ttt.cell(row, col + 1).value is not None:
                ttt.cell(row, col, 10)
    return book


def add_station_yard(book, down="[99]/100", up="[99]/100"):
    sheet = book["NODE_INFO"]
    headers = [cell.value for cell in sheet[1]]
    source = next(
        dict(zip(headers, values))
        for values in sheet.iter_rows(min_row=2, values_only=True)
        if values[headers.index("NODE_NAME")] == "TianjinNan"
    )
    lines = book["LINE_INFO"]
    line_headers = [cell.value for cell in lines[1]]
    line_row = next(
        row
        for row in range(2, lines.max_row + 1)
        if lines.cell(row, line_headers.index("LINE_IND") + 1).value == 2
    )
    line = dict(zip(line_headers, [c.value for c in lines[line_row]]))
    source.update(
        LINE_IND=2,
        LINE_CODE=line["LINE_CODE"],
        LINE_NAME_CHN=line["LINE_NAME_CHN"],
        NODE_SEQ_OF_LINE=2,
        DB_TRACKS=down,
        UB_TRACKS=up,
    )
    sheet.append([source.get(key) for key in headers])
    if "NUM_NODE" in line_headers:
        lines.cell(line_row, line_headers.index("NUM_NODE") + 1, 2)


def content(book):
    result = BytesIO()
    book.save(result)
    return result.getvalue()


def test_compact_track_fields_store_roles_separately_and_union_real_numbers():
    row = {"DB_TRACKS": "[15,10,12]/9", "UB_TRACKS": "[10,11,12]/16"}
    row.update(normalize_track_fields(row))
    down, up, numbers = member_tracks(row)
    assert down.stop_tracks == (10, 12, 15) and down.main_track == 9
    assert up.stop_tracks == (10, 11, 12) and up.main_track == 16
    assert row["DB_STOP_TRACKS"] == [10, 12, 15] and row["DB_MAIN_TRACK"] == 9
    assert row["UB_STOP_TRACKS"] == [10, 11, 12] and row["UB_MAIN_TRACK"] == 16
    assert numbers == [9, 10, 11, 12, 15, 16]
    assert parse_direction("[]/9").numbers == [9]
    assert parse_direction("[9,10]/9").numbers == [9, 10]
    with pytest.raises(ValueError, match="冲突"):
        normalize_track_fields(
            {"DB_TRACKS": "[10]/9", "UB_TRACKS": "[10]/9", "UP_TRACKS": "[9]/10"}
        )


@pytest.mark.parametrize(
    "value",
    [
        12,
        "[10,]/9",
        "[10]/",
        "[10]/0",
        "[10]/9/16",
        "[10]/[9]",
        "[true]/9",
        "[10.0]/9",
        "[0]/9",
        "[10,10]/9",
        "[10]/2147483648",
        "[10]/9.0",
    ],
)
def test_invalid_new_track_fields_are_rejected(value):
    with pytest.raises(ValueError):
        parse_direction(value)


def test_preflight_points_to_exact_track_cell_and_main_tracks_validate_by_station():
    book = track_book()
    sheet = book["NODE_INFO"]
    header = [c.value for c in sheet[1]]
    down = header.index("DB_TRACKS") + 1
    sheet.cell(2, down, "[10,11]/")
    report = parse_base(content(book), date(2025, 5, 25), "rebase", track_lists=True)
    error = next(
        i for i in report["issues"] if i["sheet"] == "NODE_INFO" and i["field"] == "DB_TRACKS"
    )
    assert (
        error["row"] == 2
        and error["node_name"] == "BeijingNan"
        and error["node_name_chn"] == "北京南"
    )
    assert error["cell"].endswith("2") and "正线" in error["message"]
    sheet.cell(2, down, "[10,11,12,13,14,15]/9")
    ttt = book["TTT1base"]
    # A main track is valid even though it is absent from the stopping-track list.
    beijing_g = next(
        c for c in range(8, ttt.max_column + 1, 4) if ttt.cell(2, c + 1).value is not None
    )
    ttt.cell(2, beijing_g, 9)
    valid = parse_base(content(book), date(2025, 5, 25), "rebase", track_lists=True)
    assert not [i for i in valid["issues"] if i["severity"] == "error"]
    # G validation uses OR across this station's lines; plotting still uses one selected line.
    add_station_yard(book)
    station_row = next(r for r in range(2, ttt.max_row + 1) if ttt.cell(r, 5).value == "TianjinNan")
    ttt.cell(station_row, 4, 1)
    first_g = next(
        c for c in range(8, ttt.max_column + 1, 4) if ttt.cell(station_row, c).value is not None
    )
    ttt.cell(station_row, first_g, 99)
    accepted = parse_base(content(book), date(2025, 5, 25), "rebase", track_lists=True)
    assert not [i for i in accepted["issues"] if i["severity"] == "error"]
    ttt.cell(station_row, first_g, 98)
    invalid = parse_base(content(book), date(2025, 5, 25), "rebase", track_lists=True)
    assert any(
        i["node_name"] == "TianjinNan" and "所属线路(1,2)" in i["message"]
        for i in invalid["issues"]
        if i["severity"] == "error"
    )
    book.close()


def resource(tracks=(9, 10, 16)):
    return {
        "node_id": "station",
        "node_name": "S",
        "node_name_chn": "车站",
        "line_ind": 1,
        "can_expand": True,
        "tracks": [{"number": t, "roles": []} for t in tracks],
        "track_count": len(tracks),
        "directions": {},
        "issues": [],
    }


def series(direction="down", track=10, arrival=100, departure=200):
    event = {
        "id": "event",
        "node_id": "station",
        "node_name": "S",
        "node_name_chn": "车站",
        "visit_seq": 2,
        "track_code": track,
        "arrival": None if arrival is None else {"seconds": arrival},
        "departure": None if departure is None else {"seconds": departure},
        "source": {"sheet": "TTT", "row": 3},
    }
    points = [
        {"event_id": "event", "field": key, "seconds": value, "y": 1000}
        for key, value in (("arrival", arrival), ("departure", departure))
        if value is not None
    ]
    return {
        "key": "run:r:t",
        "object_type": "run",
        "object_id": "r",
        "train_id": "t",
        "train_course_id": "G1",
        "train_base_date": "2025-05-25",
        "direction": direction,
        "events": [event],
        "points": points,
        "segments": [{"from": 0, "to": 1, "kind": "dwell", "anomaly": False}]
        if len(points) == 2
        else [],
    }


def test_track_occupation_uses_real_number_for_yita_and_direction_boundaries():
    for direction, arrival_level in (("down", -2), ("up", 2)):
        train = series(direction)
        actual = deepcopy(train["events"][0])
        expanded, issues = expand_track_usage([train], [resource()], ["station"], 2)
        assert expanded == ["station"] and not issues
        assert train["points"][0]["track_offset"] == arrival_level
        assert train["points"][1]["track_offset"] == -arrival_level
        assert [p["seconds"] for p in train["points"]] == [100, 200, 120, 180]
        assert train["points"][2]["track_offset"] == 0  # middle actual track 10, not 10*gap
        assert [s["kind"] for s in train["segments"]] == [
            "track_enter",
            "track_occupancy",
            "track_leave",
        ]
        assert {k: v for k, v in train["events"][0].items() if k != "track_usage"} == actual


def test_occupation_missing_tracks_times_inversions_midnight_and_zero_length():
    train = series(track=None)
    _, issues = expand_track_usage([train], [resource()], ["station"], 1)
    assert any(i["code"] == "track_missing" for i in issues)
    assert len(train["points"]) == 2 and not train["segments"]
    train = series(track=None)
    _, issues = expand_track_usage([train], [resource((1, 9, 10))], ["station"], 1)
    assert (
        train["events"][0]["track_code"] is None
        and train["events"][0]["track_usage"]["adopted_track"] == 1
    )
    assert any(i["code"] == "track_default" and i["source"]["row"] == 3 for i in issues)
    for track, code in ((99, "track_invalid"), ("10", "track_invalid")):
        train = series(track=track)
        _, issues = expand_track_usage([train], [resource()], ["station"], 1)
        assert any(i["code"] == code for i in issues) and not train["segments"]
    train = series(departure=105)
    _, issues = expand_track_usage([train], [resource()], ["station"], 1)
    assert any(i["code"] == "track_time_order" for i in issues) and len(train["points"]) == 2
    train = series(arrival=None)
    _, issues = expand_track_usage([train], [resource()], ["station"], 1)
    assert [s["kind"] for s in train["segments"]] == ["track_leave"]
    assert train["events"][0]["arrival"] is None and any(
        i["code"] == "track_missing_time" for i in issues
    )
    for arrival, departure, expected in ((86390, 86450, [86400, 86440]), (100, 120, [110, 110])):
        train = series(arrival=arrival, departure=departure)
        _, issues = expand_track_usage([train], [resource()], ["station"], 1)
        assert not issues and [p["seconds"] for p in train["points"]][2:] == expected
        assert train["events"][0]["track_usage"]["can_occupy"]
    train = series(direction="unknown")
    _, issues = expand_track_usage([train], [resource()], ["station"], 1)
    assert any(i["code"] == "track_direction" for i in issues) and not train["segments"]


def test_track_resources_do_not_mix_lines_and_support_stop_only_members():
    station = {"node_id": "station", "node_name": "S", "node_name_chn": "车站"}
    rows = [
        {"node_name": "S", "line_ind": 1, "DB_TRACKS": "[10,12]/9", "UB_TRACKS": "[12]/16"},
        {"node_name": "S", "line_ind": 2, "DB_TRACKS": "[99]/100", "UB_TRACKS": "[99]/101"},
    ]
    first = track_resources({"members": rows}, {"stations": [station]}, 1)[0]
    second = track_resources({"members": rows}, {"stations": [station]}, 2)[0]
    assert [t["number"] for t in first["tracks"]] == [9, 10, 12, 16]
    assert [t["number"] for t in second["tracks"]] == [99, 100, 101]
    assert next(t for t in first["tracks"] if t["number"] == 9)["roles"] == ["down_main"]
    rows[0]["DB_TRACKS"] = "[9,10,12]"
    stop_only = track_resources({"members": rows}, {"stations": [station]}, 1)[0]
    assert stop_only["can_expand"] and not stop_only["issues"]
    assert stop_only["directions"]["down"] == {"stop_tracks": [9, 10, 12], "main_track": None}
    rows[0]["DB_TRACKS"] = "[10,12]/9"
    duplicate = track_resources({"members": rows + [rows[0]]}, {"stations": [station]}, 1)[0]
    assert not duplicate["can_expand"]


def test_saved_fields_compact_export_and_run_base_occupation_are_read_only(tmp_path):
    book = track_book()
    source = content(book)
    book.close()
    app = create_app(Settings(database=tmp_path / "db"))
    with TestClient(app) as client:
        ids = branch(client, "tracks")
        _, base = create_base(client, ids[2], source)
        base_url = f"/api/v1/base-timetables/{base['id']}/tables"
        saved = client.get(base_url).json()
        member = saved["infrastructure"]["members"][0]
        assert member["TRACKS_SCHEMA_VERSION"] == 2
        assert member["DB_STOP_TRACKS"] == list(range(10, 16)) and member["DB_MAIN_TRACK"] == 9
        assert member["UB_MAIN_TRACK"] == 16
        node = next(n for n in saved["nodes"] if n["node_name"] == "Langfang")
        url = f"/api/v1/base-timetables/{base['id']}"
        exported = client.post(
            url + "/export",
            params={
                "revision": saved["revision"],
                "sheets": "TTT,LINE_INFO,NODE_INFO,MINIMUM_RUN_TIME",
            },
        )
        assert exported.status_code == 200
        workbook = load_workbook(BytesIO(exported.content))
        sheet = workbook["NODE_INFO"]
        headers = [c.value for c in sheet[1]]
        assert not any(
            key in headers for key in ("DB_STOP_TRACKS", "DB_MAIN_TRACK", "TRACKS_SCHEMA_VERSION")
        )
        assert sheet.cell(2, headers.index("DB_TRACKS") + 1).value == "[10,11,12,13,14,15]/9"
        workbook.close()
        parsed = parse_base(exported.content, date(2025, 5, 25), "preserve", track_lists=True)
        assert not [i for i in parsed["issues"] if i["severity"] == "error"]
        seed_result(app, client, ids[-1], base)
        saved = client.get(base_url).json()
        run_before = client.get("/api/v1/runs/result/tables").json()
        for kind, object_id, revision in (
            ("base", base["id"], saved["revision"]),
            ("run", "result", run_before["revision"]),
        ):
            body = {
                "object_type": kind,
                "object_id": object_id,
                "revision": revision,
                "line_ind": 1,
                "overlay_base": kind == "run",
            }
            collapsed = client.post("/api/v1/diagram/data", json=body).json()
            body["expanded_stations"] = [node["id"]]
            response = client.post("/api/v1/diagram/data", json=body)
            assert response.status_code == 200, response.text
            plotted = response.json()
            assert plotted["can_plot"] and plotted["expanded_stations"] == [node["id"]]
            assert plotted["axis"] == collapsed["axis"]
            r = next(r for r in plotted["track_stations"] if r["node_id"] == node["id"])
            assert r["track_count"] == 8 and [t["number"] for t in r["tracks"]] == list(
                range(9, 17)
            )
            assert any(
                s["kind"] == "track_occupancy" for t in plotted["series"] for s in t["segments"]
            )
            body["expanded_stations"] = ["wrong-node"]
            assert client.post("/api/v1/diagram/data", json=body).status_code == 422
        assert client.get(base_url).json() == saved
        assert client.get("/api/v1/runs/result/tables").json() == run_before
        # Main lines are valid editing targets; a number from another yard is not.
        event = next(e for e in run_before["events"] if e["node_id"] == node["id"])
        body = {
            "revision": run_before["revision"],
            "changes": [{"event_id": event["id"], "field": "track_code", "value": 9}],
        }
        assert client.post("/api/v1/runs/result/events/validate", json=body).json()["can_save"]
        body["changes"][0]["value"] = 99
        assert not client.post("/api/v1/runs/result/events/validate", json=body).json()["can_save"]


def test_pre_upgrade_import_preview_requires_new_preflight_without_writing_data(tmp_path):
    import json

    app = create_app(Settings(database=tmp_path / "db"))
    with TestClient(app) as client:
        ids = branch(client, "upgrade")
        book = track_book()
        response = preview(client, ids[2], content(book))
        book.close()
        assert response.status_code == 201
        job = response.json()
        assert job["report"]["tracks_schema_version"] == 2 and job["report"]["can_commit"]
        old_report = dict(job["report"])
        old_report.pop("tracks_schema_version")
        with app.state.database.transaction() as c:
            c.execute(
                "UPDATE base_tt_jobs SET report=? WHERE id=?",
                [json.dumps(old_report), job["job_id"]],
            )
        restored = client.get(f"/api/v1/ingest/jobs/{job['job_id']}").json()
        assert not restored["report"]["can_commit"]
        assert any(i["field"] == "tracks_format" for i in restored["report"]["issues"])
        assert commit(client, job).status_code == 409
        assert client.get("/api/v1/base-timetables").json() == []


def test_optional_main_track_distinguishes_no_main_empty_tracks_and_unknown_data():
    row = {"DB_TRACKS": "[12,10,11]", "UB_TRACKS": "[]"}
    row.update(normalize_track_fields(row))
    down, up, tracks = member_tracks(row)
    assert down.stop_tracks == (10, 11, 12) and down.main_track is None
    assert up.stop_tracks == () and up.main_track is None
    assert tracks == [10, 11, 12]
    assert row["DB_STOP_TRACKS"] == [10, 11, 12] and row["DB_MAIN_TRACK"] is None
    assert excel_track_field(row, "DB") == "[10,11,12]"
    assert excel_track_field(row, "UB") == "[]"
    unknown = {"DB_TRACKS": None, "UB_TRACKS": "[]"}
    unknown.update(normalize_track_fields(unknown))
    assert unknown["DB_STOP_TRACKS"] is None and excel_track_field(unknown, "DB") is None
    with pytest.raises(ValueError, match="缺失"):
        member_tracks(unknown)
    axis = {"stations": [{"node_id": "s", "node_name": "S", "node_name_chn": "车站"}]}
    empty = track_resources(
        {"members": [{"node_name": "S", "line_ind": 1, "DB_TRACKS": "[]", "UB_TRACKS": "[]"}]},
        axis,
        1,
    )[0]
    assert not empty["can_expand"] and "明确为空" in empty["issues"][0]["message"]


def test_shared_main_stop_tracks_keep_both_roles_without_duplicate_horizontal_lines():
    row = {"node_name": "S", "line_ind": 1, "DB_TRACKS": "[5,6,7]/6", "UB_TRACKS": "[5,6,8]/5"}
    row.update(normalize_track_fields(row))
    axis = {"stations": [{"node_id": "s", "node_name": "S", "node_name_chn": "南京南"}]}
    result = track_resources({"members": [row]}, axis, 1)[0]
    assert result["can_expand"] and result["track_count"] == 4 and not result["issues"]
    roles = {track["number"]: set(track["roles"]) for track in result["tracks"]}
    assert roles[5] == {"down_stop", "up_stop", "up_main"}
    assert roles[6] == {"down_stop", "down_main", "up_stop"}
    alias = {"DB_TRACKS": "[5,6]", "UP_TRACKS": "[6,5]"}
    alias.update(normalize_track_fields(alias))
    assert member_tracks(alias)[1].main_track is None
    with pytest.raises(ValueError, match="冲突"):
        normalize_track_fields({"DB_TRACKS": "[5,6]", "UB_TRACKS": "[5,6]", "UP_TRACKS": "[5,6]/6"})


def test_stop_only_termini_and_shared_nanjing_tracks_roundtrip_import_edit_and_plot(tmp_path):
    import json

    from test_run_import import commit as commit_run
    from test_run_import import preview as preview_run

    book = track_book()
    sheet = book["NODE_INFO"]
    headers = [c.value for c in sheet[1]]
    for row in range(2, sheet.max_row + 1):
        name = sheet.cell(row, headers.index("NODE_NAME_CHN") + 1).value
        if name in ("北京南", "上海虹桥"):
            sheet.cell(row, headers.index("DB_TRACKS") + 1, "[10,11,12,13,14,15]")
            sheet.cell(row, headers.index("UB_TRACKS") + 1, "[10,11,12,13,14,15]")
        elif name == "南京南":
            sheet.cell(row, headers.index("DB_TRACKS") + 1, "[5,6,7]/6")
            sheet.cell(row, headers.index("UB_TRACKS") + 1, "[5,6,8]/5")
    ttt = book["TTT1base"]
    for row in range(2, ttt.max_row + 1):
        if ttt.cell(row, 1).value == "南京南":
            for col in range(8, ttt.max_column + 1, 4):
                if ttt.cell(row, col + 1).value is not None:
                    ttt.cell(row, col, 5)
    source = content(book)
    book.close()
    app = create_app(Settings(database=tmp_path / "db"))
    with TestClient(app) as client:
        ids = branch(client, "optional-main")
        _, base = create_base(client, ids[2], source)
        run_job = preview_run(client, ids[-1], source).json()
        assert run_job["report"]["can_commit"], run_job["report"]["issues"]
        run_response = commit_run(client, run_job)
        assert run_response.status_code == 200, run_response.text
        run = run_response.json()
        base_url = f"/api/v1/base-timetables/{base['id']}"
        saved = client.get(base_url + "/tables").json()
        members = saved["infrastructure"]["members"]
        for member in members:
            if member["NODE_NAME_CHN"] in ("北京南", "上海虹桥"):
                assert member["DB_MAIN_TRACK"] is None and member["UB_MAIN_TRACK"] is None
                assert member["DB_STOP_TRACKS"] == list(range(10, 16))
        with app.state.database.transaction() as c:
            persisted = json.loads(
                c.execute(
                    "SELECT document FROM base_timetables WHERE id=?", [base["id"]]
                ).fetchone()[0]
            )
        terminal = next(m for m in persisted["members"] if m["NODE_NAME_CHN"] == "北京南")
        assert terminal["DB_MAIN_TRACK"] is None and terminal["DB_STOP_TRACKS"] == list(
            range(10, 16)
        )
        exported = client.post(
            base_url + "/export",
            params={
                "revision": saved["revision"],
                "sheets": "TTT,LINE_INFO,NODE_INFO,MINIMUM_RUN_TIME",
            },
        )
        assert exported.status_code == 200, exported.text
        workbook = load_workbook(BytesIO(exported.content))
        rows = list(workbook["NODE_INFO"].values)
        values = [dict(zip(rows[0], row)) for row in rows[1:]]
        assert (
            next(m for m in values if m["NODE_NAME_CHN"] == "北京南")["DB_TRACKS"]
            == "[10,11,12,13,14,15]"
        )
        assert next(m for m in values if m["NODE_NAME_CHN"] == "南京南")["DB_TRACKS"] == "[5,6,7]/6"
        workbook.close()
        parsed = parse_base(exported.content, date(2025, 5, 25), "preserve", track_lists=True)
        assert not [i for i in parsed["issues"] if i["severity"] == "error"]
        targets = {
            n["node_name_chn"]: n["id"]
            for n in saved["nodes"]
            if n["node_name_chn"] in ("北京南", "上海虹桥", "南京南")
        }
        assert len(targets) == 3
        for kind, object_id, url in (
            ("base", base["id"], base_url),
            ("run", run["id"], f"/api/v1/runs/{run['id']}"),
        ):
            before = client.get(url + "/tables").json()
            plotted = client.post(
                "/api/v1/diagram/data",
                json={
                    "object_type": kind,
                    "object_id": object_id,
                    "revision": before["revision"],
                    "line_ind": 1,
                    "overlay_base": kind == "run",
                    "expanded_stations": list(targets.values()),
                },
            ).json()
            assert plotted["can_plot"] and set(plotted["expanded_stations"]) == set(
                targets.values()
            )
            resources = {r["node_id"]: r for r in plotted["track_stations"]}
            for name in ("北京南", "上海虹桥"):
                resource = resources[targets[name]]
                assert (
                    resource["track_count"] == 6
                    and resource["directions"]["down"]["main_track"] is None
                )
                assert all(
                    not role.endswith("_main")
                    for track in resource["tracks"]
                    for role in track["roles"]
                )
            assert resources[targets["南京南"]]["track_count"] == 4
            for name, track in (("北京南", 12), ("上海虹桥", 12), ("南京南", 6)):
                event = next(e for e in before["events"] if e["node_id"] == targets[name])
                result = client.post(
                    url + "/events/validate",
                    json={
                        "revision": before["revision"],
                        "changes": [
                            {"event_id": event["id"], "field": "track_code", "value": track}
                        ],
                    },
                ).json()
                assert result["can_save"], result["issues"]
            assert client.get(url + "/tables").json() == before
