from datetime import date
from io import BytesIO

from app.core.config import Settings
from app.ingest.base_parser import parse_base
from app.main import create_app
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sample_workbook import SAMPLE_WORKBOOK
from test_base_import import create_base, preview
from test_catalog_deletion import branch


def test_invalid_source_row_beyond_preview_keeps_raw_node_name_and_index():
    book = load_workbook(BytesIO(SAMPLE_WORKBOOK.read_bytes()))
    sheet = book["NODE_INFO"]
    headers = [cell.value for cell in sheet[1]]
    original = next(
        list(row) for row in sheet.values if row[headers.index("NODE_NAME")] == "Langfang"
    )
    padding = original.copy()
    original[headers.index("NODE_IND")] = 0
    # Preview counts nonempty records, so use repeated source rows rather than blank padding.
    while sheet.max_row < 250:
        sheet.append(padding)
    sheet.append(original)
    data = BytesIO()
    book.save(data)
    book.close()
    report = parse_base(data.getvalue(), date(2025, 5, 25), "rebase", track_lists=True)
    error = next(i for i in report["issues"] if i["sheet"] == "NODE_INFO" and i["row"] == 251)
    assert error["severity"] == "error"
    assert error["node_ind"] == 0
    assert error["node_name"] == "Langfang"
    assert error["node_name_chn"] == "廊坊"
    table = next(t for t in report["previews"] if t["name"] == "NODE_INFO")
    assert all(row["row"] != 251 for row in table["rows"])
    # The valid occurrence of this node elsewhere must not overwrite the invalid source value.
    assert next(n for n in report["nodes"] if n["node_name"] == "Langfang")["node_ind"] == 2


def test_ttt_cell_error_has_node_labels_but_sheet_errors_do_not_invent_them():
    book = load_workbook(BytesIO(SAMPLE_WORKBOOK.read_bytes()))
    book["TTT1base"]["K2"] = "bad time"
    del book["MINIMUM_RUN_TIME"]
    data = BytesIO()
    book.save(data)
    book.close()
    report = parse_base(data.getvalue(), date(2025, 5, 25), "rebase", track_lists=True)
    error = next(i for i in report["issues"] if i["sheet"] == "TTT1base" and i["cell"] == "K2")
    assert error["node_ind"] == 1 and error["node_name"] == "BeijingNan"
    assert error["node_name_chn"] == "北京南"
    missing = next(
        i for i in report["issues"] if i["sheet"] == "MINIMUM_RUN_TIME" and i["field"] == "Sheet"
    )
    assert "node_name" not in missing and "node_ind" not in missing


def test_network_index_conflicts_return_exact_node_and_source_in_restored_preview(tmp_path):
    app = create_app(Settings(database=tmp_path / "db"))
    with TestClient(app) as client:
        ids = branch(client, "labels")
        _, base = create_base(client, ids[2], SAMPLE_WORKBOOK.read_bytes())
        with app.state.database.transaction() as c:
            node_id = c.execute("SELECT id FROM nodes WHERE node_name='Langfang'").fetchone()[0]
            c.execute("INSERT INTO network_node_maps VALUES (?,?,?)", [ids[1], 2, node_id])
        book = load_workbook(BytesIO(SAMPLE_WORKBOOK.read_bytes()))
        for title in ("NODE_INFO", "TTT1base"):
            sheet = book[title]
            headers = [cell.value for cell in sheet[1]]
            for row in range(2, sheet.max_row + 1):
                if sheet.cell(row, headers.index("NODE_NAME") + 1).value == "Langfang":
                    sheet.cell(row, headers.index("NODE_IND") + 1, 999)
        data = BytesIO()
        book.save(data)
        book.close()
        response = preview(client, ids[-1], data.getvalue())
        assert response.status_code == 201, response.text
        job = response.json()
        assert not job["report"]["can_commit"]
        errors = [
            i
            for i in job["report"]["issues"]
            if i["severity"] == "error" and i["field"] == "NODE_IND"
        ]
        assert len(errors) >= 2  # Both registered network mapping and other BaseTT scope checks.
        assert all(i["node_ind"] == 999 and i["node_name"] == "Langfang" for i in errors)
        assert all(
            i["node_name_chn"] == "廊坊" and i["sheet"] == "TTT1base" and i["row"] == 3
            for i in errors
        )
        restored = client.get(f"/api/v1/ingest/jobs/{job['job_id']}").json()
        assert restored["report"]["issues"] == job["report"]["issues"]
        assert client.get(f"/api/v1/base-timetables/{base['id']}").json()["revision"] == 1
