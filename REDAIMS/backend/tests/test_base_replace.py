from datetime import date
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from app.core.config import Settings
from app.ingest.base_parser import parse_base
from app.ingest.base_service import BaseImportService, BindingChange, CommitImport, ImportOptions
from app.main import create_app
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sample_workbook import SAMPLE_WORKBOOK
from test_base_import import commit, create_base, preview
from test_catalog_deletion import branch


@pytest.fixture
def sample():
    return SAMPLE_WORKBOOK.read_bytes()


def test_editor_validation_save_stale_and_cross_day(tmp_path, sample):
    with TestClient(create_app(Settings(database=tmp_path / "db"))) as client:
        ids = branch(client, "edit")
        _, base = create_base(client, ids[-1], sample)
        url = f"/api/v1/base-timetables/{base['id']}"
        before = client.get(url + "/tables").json()
        e = next(e for e in before["events"] if e["departure"] is not None)
        changes = [{"event_id": e["id"], "field": "departure", "value": "25:10:00"}]
        payload = {"revision": before["revision"], "changes": changes}
        response = client.post(url + "/events/validate", json=payload)
        assert response.status_code == 200, response.text
        valid = response.json()
        assert valid["can_save"] and valid["train_date_changes"]
        assert client.get(url + "/tables").json()["events"] == before["events"]
        assert client.post(url + "/events/edit", json=payload).status_code == 422
        save = payload | {"validation_token": valid["validation_token"], "warnings_confirmed": True}
        assert client.post(url + "/events/edit", json=save).status_code == 422
        saved = client.post(url + "/events/edit", json=save | {"train_dates_confirmed": True})
        assert saved.status_code == 200, saved.text
        after = saved.json()
        changed = next(x for x in after["events"] if x["id"] == e["id"])
        assert changed["departure"] == {"hms": "01:10:00", "seconds": 90600, "day_offset": 1}
        assert [x for x in after["events"] if x["id"] != e["id"]] == [
            x for x in before["events"] if x["id"] != e["id"]
        ]
        assert Path(base["source_files"][0]["managed_path"]).read_bytes() == sample
        assert (
            client.post(
                url + "/events/edit", json=save | {"train_dates_confirmed": True}
            ).status_code
            == 409
        )


def test_editor_errors_foreign_cells_and_changed_validation(tmp_path, sample):
    with TestClient(create_app(Settings(database=tmp_path / "db"))) as client:
        ids = branch(client, "edit")
        _, base = create_base(client, ids[-1], sample)
        url = f"/api/v1/base-timetables/{base['id']}"
        before = client.get(url + "/tables").json()
        e = before["events"][0]
        for change in [
            {"event_id": e["id"], "field": "arrival", "value": "bad"},
            {"event_id": e["id"], "field": "visit_seq", "value": 0},
            {"event_id": str(uuid4()), "field": "arrival", "value": "08:00:00"},
        ]:
            payload = {"revision": before["revision"], "changes": [change]}
            report = client.post(url + "/events/validate", json=payload).json()
            assert not report["can_save"]
            assert (
                client.post(
                    url + "/events/edit",
                    json=payload
                    | {
                        "validation_token": report["validation_token"],
                        "warnings_confirmed": True,
                        "train_dates_confirmed": True,
                    },
                ).status_code
                == 422
            )
        payload = {
            "revision": before["revision"],
            "changes": [{"event_id": e["id"], "field": "arrival", "value": "08:00:00"}],
        }
        report = client.post(url + "/events/validate", json=payload).json()
        payload["changes"][0]["value"] = "08:01:00"
        assert (
            client.post(
                url + "/events/edit",
                json=payload
                | {"validation_token": report["validation_token"], "warnings_confirmed": True},
            ).status_code
            == 409
        )
        assert client.get(url + "/tables").json()["events"] == before["events"]


def test_cover_failure_and_interrupted_recovery(tmp_path, sample, monkeypatch):
    import app.base_timetables.service as module

    settings = Settings(database=tmp_path / "db")
    with TestClient(create_app(settings)) as client:
        ids = branch(client, "x")
        _, base = create_base(client, ids[-1], sample)
        old_file = Path(base["source_files"][0]["managed_path"])
        job = preview(client, ids[-1], sample, "changed").json()
        original = module.shutil.copyfile

        def fail(*args):
            raise OSError("disk error")

        monkeypatch.setattr(module.shutil, "copyfile", fail)
        assert commit(client, job, replace_confirmed=True).status_code == 503
        assert client.get(f"/api/v1/base-timetables/{base['id']}").json()["name"] == base["name"]
        assert old_file.read_bytes() == sample
        monkeypatch.setattr(module.shutil, "copyfile", original)
        assert commit(client, job, replace_confirmed=True).status_code == 200
    # Recovery removes only an uncommitted new source, not a committed base.
    app = create_app(settings)
    with TestClient(app) as client:
        pending = preview(client, ids[-1], sample).json()
        with app.state.database.transaction() as c:
            c.execute("UPDATE base_tt_jobs SET state='prepared' WHERE id=?", [pending["job_id"]])
        partial = tmp_path / "source_files" / (pending["job_id"] + ".xlsx")
        partial.write_bytes(b"partial")
    with TestClient(create_app(settings)) as client:
        assert not partial.exists()
        assert client.get(f"/api/v1/ingest/jobs/{pending['job_id']}").json()["state"] == "failed"
        assert len(client.get("/api/v1/base-timetables").json()) == 1


def test_explicit_migration_preserves_events_sources_and_backup(tmp_path, sample):
    settings = Settings(database=tmp_path / "db")
    app = create_app(settings)
    with TestClient(app) as client:
        ids = branch(client, "legacy")
        other = branch(client, "other")
        legacy = BaseImportService(app.state.database, settings)
        job = legacy.preview(
            sample,
            "old.xlsx",
            ImportOptions(batch_id=ids[-1], name="原基本图", date_base=date(2025, 5, 25)),
        )
        old = legacy.commit(
            job["job_id"],
            CommitImport(
                confirmed=True,
                date_confirmed=True,
                train_dates_confirmed=True,
                warnings_confirmed=True,
            ),
        )
        binding = legacy.binding_preview(other[3])
        legacy.bind(
            other[3],
            BindingChange(
                base_run_id=old["id"], confirmed=True, preview_token=binding["preview_token"]
            ),
        )
        assert client.get("/api/v1/base-timetables").json() == []
        assert client.get("/api/v1/runs").json() == []
        assert len(client.get("/api/v1/base-timetable-migrations").json()) == 1
        route = f"/api/v1/base-timetable-migrations/{old['id']}"
        plan = client.post(route + "/preview", json={"target_category_id": ids[3]}).json()
        assert plan["can_migrate"]
        assert any(r["after"]["mode"] == "missing" for r in plan["legacy_binding_changes"])
        with app.state.database.transaction() as c:
            assert c.execute("SELECT count(*) FROM runs").fetchone()[0] == 1
        response = client.post(
            route + "/commit",
            json={
                "target_category_id": ids[3],
                "confirmed": True,
                "preview_token": plan["preview_token"],
            },
        )
        assert response.status_code == 200, response.text
        result = response.json()
        assert Path(result["backup_path"], "db").is_file()
        assert Path(result["backup_path"], "source_files").is_dir()
        assert result["base"]["summary"]["events"] == 148
        assert result["base"]["source_files"] == old["source_files"]
        assert (
            len(
                client.get(f"/api/v1/base-timetables/{result['base']['id']}/tables").json()[
                    "events"
                ]
            )
            == 148
        )
        assert client.get("/api/v1/base-timetable-migrations").json() == []
        with app.state.database.transaction() as c:
            assert c.execute("SELECT count(*) FROM runs").fetchone()[0] == 0
            assert c.execute("SELECT count(*) FROM base_bindings").fetchone()[0] == 0
            assert c.execute("SELECT count(*) FROM base_tt_migrations").fetchone()[0] == 1


def test_migration_occupied_missing_source_and_stale_preview(tmp_path, sample):
    settings = Settings(database=tmp_path / "db")
    app = create_app(settings)
    with TestClient(app) as client:
        ids = branch(client, "old")
        legacy = BaseImportService(app.state.database, settings)
        job = legacy.preview(
            sample,
            "old.xlsx",
            ImportOptions(batch_id=ids[-1], name="old", date_base=date(2025, 5, 25)),
        )
        old = legacy.commit(
            job["job_id"],
            CommitImport(
                confirmed=True,
                date_confirmed=True,
                train_dates_confirmed=True,
                warnings_confirmed=True,
            ),
        )
        url = f"/api/v1/base-timetable-migrations/{old['id']}"
        plan = client.post(url + "/preview", json={"target_category_id": ids[3]}).json()
        create_base(client, ids[3], sample)
        assert (
            client.post(
                url + "/commit",
                json={
                    "target_category_id": ids[3],
                    "confirmed": True,
                    "preview_token": plan["preview_token"],
                },
            ).status_code
            == 409
        )
        assert not client.post(url + "/preview", json={"target_category_id": ids[3]}).json()[
            "can_migrate"
        ]
        Path(old["source_files"][0]["managed_path"]).write_bytes(b"modified")
        assert not client.post(url + "/preview", json={"target_category_id": ids[2]}).json()[
            "can_migrate"
        ]
        assert len(client.get("/api/v1/base-timetable-migrations").json()) == 1


def test_track_lists_alias_and_shared_tracks(sample):
    workbook = load_workbook(BytesIO(sample))
    sheet = workbook["NODE_INFO"]
    headers = [cell.value for cell in sheet[1]]
    down, up = headers.index("DB_TRACKS") + 1, headers.index("UB_TRACKS") + 1
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, down, "[1,2,6,8,17]/9")
        sheet.cell(row, up, "[1,2,6,8,17]/9")
    sheet.cell(1, up, "UP_TRACKS")
    buffer = BytesIO()
    workbook.save(buffer)
    report = parse_base(buffer.getvalue(), date(2025, 5, 25), "rebase", track_lists=True)
    assert all(isinstance(m["DB_TRACKS"], str) for m in report["members"])
    assert not any(i["field"] == "UB_TRACKS" and i["severity"] == "error" for i in report["issues"])
    sheet.cell(2, down, 12)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    invalid = parse_base(buffer.getvalue(), date(2025, 5, 25), "rebase", track_lists=True)
    assert any("列表字符串" in i["message"] for i in invalid["issues"])


def seed_result(app, client, batch, base, run_id="result"):
    """A real result dataset, without using base-run storage or an import alias."""
    from app.ingest.base_service import encode

    doc = client.get(f"/api/v1/base-timetables/{base['id']}/tables").json()
    node_names = {n["id"]: n["node_name_chn"] for n in doc["nodes"]}
    trains = {t["id"]: str(uuid4()) for t in doc["trains"]}
    with app.state.database.transaction() as c:
        c.execute("INSERT INTO runs(id,batch_id,name) VALUES (?,?,?)", [run_id, batch, "结果"])
        c.execute("INSERT INTO run_working_data VALUES (?,?,1)", [run_id, base["date_base"]])
        for t in doc["trains"]:
            c.execute(
                "INSERT INTO run_trains VALUES (?,?,?,?,?)",
                [trains[t["id"]], run_id, t["train_course_id"], t["train_base_date"], encode(t)],
            )
        for e in doc["events"]:
            a, d = e["arrival"] or {}, e["departure"] or {}
            c.execute(
                "INSERT INTO timetable_events VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                [
                    str(uuid4()),
                    run_id,
                    trains[e["train_id"]],
                    e["node_id"],
                    e["train_course_id"],
                    e["node_name"],
                    node_names[e["node_id"]],
                    e["visit_seq"],
                    a.get("hms"),
                    a.get("seconds"),
                    a.get("day_offset"),
                    d.get("hms"),
                    d.get("seconds"),
                    d.get("day_offset"),
                    e["track_code"],
                    encode(e["source"]),
                    encode(e["line_inds"]),
                ],
            )
    return doc


def test_result_editor_isolation_and_batch_reclassification(tmp_path, sample):
    app = create_app(Settings(database=tmp_path / "db"))
    with TestClient(app) as client:
        ids = branch(client, "run")
        _, base = create_base(client, ids[3], sample)
        base_before = seed_result(app, client, ids[-1], base)
        url = "/api/v1/runs/result"
        before = client.get(url + "/tables").json()
        event = next(e for e in before["events"] if e["departure"] is not None)
        changes = [{"event_id": event["id"], "field": "departure", "value": "11:11:11"}]
        payload = {"revision": before["revision"], "changes": changes}
        valid = client.post(url + "/events/validate", json=payload).json()
        assert valid["can_save"], valid
        response = client.post(
            url + "/events/edit",
            json=payload
            | {
                "validation_token": valid["validation_token"],
                "warnings_confirmed": True,
                "train_dates_confirmed": True,
            },
        )
        assert response.status_code == 200, response.text
        after = response.json()
        assert (
            next(e for e in after["events"] if e["id"] == event["id"])["departure"]["seconds"]
            == 40271
        )
        assert [e for e in before["events"] if e["id"] != event["id"]] == [
            e for e in after["events"] if e["id"] != event["id"]
        ]
        assert (
            client.get(f"/api/v1/base-timetables/{base['id']}/tables").json()["events"]
            == base_before["events"]
        )
        assert client.post(url + "/events/validate", json=payload).status_code == 409
        # Shared editor cannot move a result to an arbitrary category.
        assert (
            client.post(url + "/move-preview", json={"target_category_id": ids[3]}).status_code
            == 422
        )
        batch = client.post(
            "/api/v1/catalog", json={"kind": "batch", "name": "目标批次", "parent_id": ids[4]}
        ).json()["id"]
        plan = client.post(url + "/move-preview", json={"target_category_id": batch}).json()
        moved = client.post(
            url + "/move",
            json={
                "target_category_id": batch,
                "preview_token": plan["preview_token"],
                "confirmed": True,
            },
        )
        assert moved.status_code == 200, moved.text
        assert moved.json()["owner_category_id"] == batch
        assert moved.json()["effective_base"]["base_tt_id"] == base["id"]
        # Validation must also compare ordinary results against independent bases.
        fresh = client.get(url + "/tables").json()
        train = next(
            t for t in fresh["trains"] if sum(e["train_id"] == t["id"] for e in fresh["events"]) > 2
        )
        visits = sorted(
            [e for e in fresh["events"] if e["train_id"] == train["id"]],
            key=lambda e: e["visit_seq"],
        )
        swapped = [
            {"event_id": visits[i]["id"], "field": "visit_seq", "value": visits[3 - i]["visit_seq"]}
            for i in (1, 2)
        ]
        report = client.post(
            url + "/events/validate", json={"revision": fresh["revision"], "changes": swapped}
        ).json()
        assert not report["can_save"] and any("SEQ" in i["message"] for i in report["issues"])


def test_import_checks_result_sequences_and_delete_cleans_metadata(tmp_path, sample):
    from test_catalog_deletion import execute

    app = create_app(Settings(database=tmp_path / "db"))
    with TestClient(app) as client:
        ids = branch(client, "run")
        _, base = create_base(client, ids[3], sample)
        seed_result(app, client, ids[-1], base)
        with app.state.database.transaction() as c:
            c.execute("UPDATE timetable_events SET visit_seq=visit_seq+100 WHERE run_id='result'")
        report = preview(client, ids[2], sample).json()["report"]
        assert not report["can_commit"]
        assert any("普通 run" in i["message"] for i in report["issues"])
        plan = client.get(f"/api/v1/catalog/{ids[-1]}/deletion-preview").json()
        response = execute(client, ids[-1], plan, "delete")
        assert response.status_code == 200, response.text
        with app.state.database.transaction() as c:
            assert c.execute("SELECT count(*) FROM run_working_data").fetchone()[0] == 0
        assert client.get(f"/api/v1/base-timetables/{base['id']}").status_code == 200


def test_replace_database_failure_rolls_back_every_table(tmp_path, sample, monkeypatch):
    from contextlib import contextmanager

    app = create_app(Settings(database=tmp_path / "db"))
    with TestClient(app) as client:
        ids = branch(client, "atomic")
        original_job, base = create_base(client, ids[-1], sample)
        old = client.get(f"/api/v1/base-timetables/{base['id']}/tables").json()
        job = preview(client, ids[-1], sample, "替换失败").json()
        database = app.state.database
        transaction = database.transaction

        class FailingConnection:
            def __init__(self, c):
                self.c = c

            def execute(self, sql, *args):
                if sql.startswith("INSERT INTO base_tt_sources"):
                    raise RuntimeError("injected transaction failure")
                return self.c.execute(sql, *args)

        @contextmanager
        def fail_transaction():
            with transaction() as c:
                yield FailingConnection(c)

        monkeypatch.setattr(database, "transaction", fail_transaction)
        with pytest.raises(RuntimeError, match="injected transaction"):
            commit(client, job, replace_confirmed=True)
        monkeypatch.setattr(database, "transaction", transaction)
        assert client.get(f"/api/v1/base-timetables/{base['id']}/tables").json() == old
        assert Path(base["source_files"][0]["managed_path"]).read_bytes() == sample
        assert not (tmp_path / "source_files" / (job["job_id"] + ".xlsx")).exists()
        assert (
            client.get(f"/api/v1/ingest/jobs/{original_job['job_id']}").json()["state"]
            == "committed"
        )
        with transaction() as c:
            assert c.execute("SELECT count(*) FROM file_cleanup").fetchone()[0] == 0
