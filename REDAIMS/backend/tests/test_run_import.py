from io import BytesIO
from pathlib import Path

import pytest
from app.core.config import Settings
from app.main import create_app
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sample_workbook import SAMPLE_WORKBOOK
from test_base_import import confirmation, create_base
from test_catalog_deletion import branch


def result_book(change=None, poison=False, title="TTT1base"):
    book = load_workbook(BytesIO(SAMPLE_WORKBOOK.read_bytes()))
    sheet = book["TTT1base"]
    sheet.title = title
    if change:
        change(sheet)
    for name in list(book.sheetnames):
        if name != title:
            if poison:
                other = book[name]
                other.delete_rows(1, other.max_row)
                other.append(["wrong header", "=1/0", "invalid foundation data"])
            else:
                del book[name]
    data = BytesIO()
    book.save(data)
    book.close()
    return data.getvalue()


def preview(client, batch, content=None, **params):
    return client.post(
        "/api/v1/runs/import/preview",
        params={
            "batch_id": batch,
            "name": "调度结果",
            "filename": "result.xlsx",
            "date_base": "2025-05-25",
            **params,
        },
        content=result_book() if content is None else content,
    )


def commit(client, job, **extra):
    return client.post(
        f"/api/v1/runs/import/{job['job_id']}/commit", json=confirmation(job, **extra)
    )


def imported(client, batch, **params):
    response = preview(client, batch, **params)
    assert response.status_code == 201, response.text
    job = response.json()
    assert job["report"]["can_commit"], job["report"]["issues"]
    response = commit(client, job)
    assert response.status_code == 200, response.text
    return job, response.json()


def test_ttt_only_import_inherits_base_and_opens_full_shared_editor_and_export(tmp_path):
    settings = Settings(database=tmp_path / "db")
    app = create_app(settings)
    with TestClient(app) as client:
        ids = branch(client, "run")
        _, base = create_base(client, ids[2], SAMPLE_WORKBOOK.read_bytes())
        base_url = f"/api/v1/base-timetables/{base['id']}/tables"
        base_before = client.get(base_url).json()
        content = result_book(poison=True)
        response = preview(client, ids[-1], content)
        assert response.status_code == 201, response.text
        job = response.json()
        assert job["report"]["can_commit"], job["report"]["issues"]
        assert job["report"]["effective_base"]["base_tt_id"] == base["id"]
        assert {p["name"] for p in job["report"]["previews"]} == {"TTT1base"}
        assert job["report"]["lines"] == job["report"]["members"] == []
        saved = commit(client, job)
        assert saved.status_code == 200, saved.text
        run = saved.json()
        assert run["run_no"] == 1 and run["owner_category_id"] == ids[-1]
        assert commit(client, job).json()["id"] == run["id"]
        assert client.get(base_url).json()["events"] == base_before["events"]
        tables = client.get(f"/api/v1/runs/{run['id']}/tables").json()
        assert len(tables["events"]) == 148 and len(tables["trains"]) == 10
        assert len(tables["route_nodes"]) > len(tables["nodes"])
        assert tables["display_lines"] == base_before["display_lines"]
        assert all(n["node_ind"] and n["line_inds"] for n in tables["nodes"])
        with app.state.database.transaction() as c:
            assert (
                c.execute(
                    "SELECT count(*) FROM base_run_data WHERE run_id=?", [run["id"]]
                ).fetchone()[0]
                == 0
            )
            assert (
                c.execute(
                    "SELECT count(*) FROM infrastructure_records WHERE run_id=?", [run["id"]]
                ).fetchone()[0]
                == 0
            )
            source = c.execute(
                "SELECT id,managed_name FROM source_files WHERE run_id=?", [run["id"]]
            ).fetchone()
        assert client.get(f"/api/v1/source-files/{source[0]}").content == content
        assert (tmp_path / "source_files" / source[1]).is_file()
        choices = client.get(f"/api/v1/runs/{run['id']}/export-options").json()
        assert choices["defaults"] == ["TTT"]
        exported = client.post(
            f"/api/v1/runs/{run['id']}/export",
            params={"revision": tables["revision"], "sheets": "TTT"},
        )
        assert exported.status_code == 200, exported.text
        book = load_workbook(BytesIO(exported.content))
        assert book.sheetnames == ["TTT1base"]
        book.close()
        assert (
            client.get(f"/api/v1/runs/import/jobs/{job['job_id']}").json()["state"] == "committed"
        )
    with TestClient(create_app(settings)) as client:
        assert client.get(f"/api/v1/runs/{run['id']}/tables").json()["events"] == tables["events"]


def test_no_base_manual_sheet_and_zero_number_remain_ordinary_with_explicit_warnings(tmp_path):
    app = create_app(Settings(database=tmp_path / "db"))
    with TestClient(app) as client:
        ids = branch(client, "plain")
        content = result_book(title="运行结果")
        auto = preview(client, ids[-1], content).json()
        assert not auto["report"]["can_commit"]
        response = preview(client, ids[-1], content, sheet="运行结果", run_no=0)
        job = response.json()
        assert job["report"]["can_commit"] and job["report"]["effective_base"]["mode"] == "missing"
        assert commit(client, job, warnings_confirmed=False).status_code == 422
        run = commit(client, job).json()
        assert run["run_no"] == 0 and run["object_type"] == "run"
        assert client.get("/api/v1/base-timetables").json() == []
        with app.state.database.transaction() as c:
            assert c.execute("SELECT count(*) FROM base_run_data").fetchone()[0] == 0
        assert client.get("/api/v1/runs/import/jobs").json()[0]["job_id"] == job["job_id"]
        assert client.get(f"/api/v1/ingest/jobs/{job['job_id']}").status_code == 404
        assert client.delete(f"/api/v1/runs/import/jobs/{auto['job_id']}").status_code == 200


def test_duplicate_auto_number_and_same_batch_target_restrictions(tmp_path):
    with TestClient(create_app(Settings(database=tmp_path / "db"))) as client:
        ids = branch(client, "ids")
        _, run = imported(client, ids[-1])
        duplicate = preview(client, ids[-1], run_no=1).json()
        assert not duplicate["report"]["can_commit"]
        assert commit(client, duplicate).status_code == 409
        _, second = imported(client, ids[-1])
        assert second["run_no"] == 2
        for kwargs in [
            {"write_mode": "replace", "run_no": 1},
            {"write_mode": "replace", "run_no": 1, "target_run_id": run["id"], "name": "不同名"},
            {"write_mode": "create", "target_run_id": run["id"]},
            {"run_no": -1},
        ]:
            assert preview(client, ids[-1], **kwargs).status_code == 422
        assert preview(client, ids[2]).status_code == 422
        assert preview(client, ids[-1], b"bad").status_code == 422
        assert preview(client, ids[-1], name=" ").status_code == 422


def test_replace_preserves_identity_requires_confirmation_and_rejects_changed_working_data(
    tmp_path,
):
    app = create_app(Settings(database=tmp_path / "db"))
    with TestClient(app) as client:
        ids = branch(client, "replace")
        old_job, run = imported(client, ids[-1])
        url = f"/api/v1/runs/{run['id']}"
        before = client.get(url + "/tables").json()
        options = {"write_mode": "replace", "run_no": run["run_no"], "target_run_id": run["id"]}
        stale = preview(client, ids[-1], **options).json()
        event = next(e for e in before["events"] if e["departure"] is not None)
        changes = [{"event_id": event["id"], "field": "departure", "value": "25:10:00"}]
        body = {"revision": before["revision"], "changes": changes}
        validation = client.post(url + "/events/validate", json=body).json()
        assert validation["can_save"], validation
        changed = client.post(
            url + "/events/edit",
            json=body
            | {
                "validation_token": validation["validation_token"],
                "warnings_confirmed": True,
                "train_dates_confirmed": True,
            },
        )
        assert changed.status_code == 200, changed.text
        assert commit(client, stale, replace_confirmed=True).status_code == 409
        assert client.get(url + "/tables").json()["revision"] == changed.json()["revision"]
        fresh = preview(client, ids[-1], **options).json()
        assert fresh["report"]["can_commit"]
        assert commit(client, fresh).status_code == 422
        result = commit(client, fresh, replace_confirmed=True)
        assert result.status_code == 200, result.text
        assert result.json()["id"] == run["id"] and result.json()["run_no"] == run["run_no"]
        after = client.get(url + "/tables").json()
        assert after["revision"] != changed.json()["revision"]
        assert (
            client.get(f"/api/v1/runs/import/jobs/{old_job['job_id']}").json()["state"]
            == "superseded"
        )
        assert len(list((tmp_path / "source_files").glob("*.xlsx"))) == 1
        assert client.delete(f"/api/v1/runs/import/jobs/{fresh['job_id']}").status_code == 409


def test_file_and_transaction_failure_leave_original_run_and_source_unchanged(
    tmp_path, monkeypatch
):
    import app.ingest.run_service as module

    app = create_app(Settings(database=tmp_path / "db"))
    with TestClient(app) as client:
        ids = branch(client, "failure")
        _, run = imported(client, ids[-1])
        url = f"/api/v1/runs/{run['id']}/tables"
        before = client.get(url).json()
        old_files = {p.name: p.read_bytes() for p in (tmp_path / "source_files").glob("*.xlsx")}
        job = preview(
            client, ids[-1], write_mode="replace", target_run_id=run["id"], run_no=1
        ).json()
        original_copy = module.shutil.copyfile

        def fail_copy(source, target):
            Path(target).write_bytes(b"partial")
            raise OSError("simulated disk failure")

        monkeypatch.setattr(module.shutil, "copyfile", fail_copy)
        assert commit(client, job, replace_confirmed=True).status_code == 503
        monkeypatch.setattr(module.shutil, "copyfile", original_copy)
        original_insert = module.RunImportService.insert_document

        def fail_transaction(self, c, *args):
            original_insert(self, c, *args)
            raise RuntimeError("simulated transaction failure")

        monkeypatch.setattr(module.RunImportService, "insert_document", fail_transaction)
        with pytest.raises(RuntimeError, match="simulated transaction failure"):
            commit(client, job, replace_confirmed=True)
        assert client.get(url).json() == before
        assert {
            p.name: p.read_bytes() for p in (tmp_path / "source_files").glob("*.xlsx")
        } == old_files
        assert client.get(f"/api/v1/runs/import/jobs/{job['job_id']}").json()["state"] == "preview"
        monkeypatch.setattr(module.RunImportService, "insert_document", original_insert)
        assert commit(client, job, replace_confirmed=True).status_code == 200


def test_identity_tracks_seq_and_cross_network_category_move_are_checked(tmp_path):
    app = create_app(Settings(database=tmp_path / "db"))
    with TestClient(app) as client:
        ids = branch(client, "check")
        # Give this test explicit known tracks; the general fixture permits unknown lists.
        book = load_workbook(BytesIO(SAMPLE_WORKBOOK.read_bytes()))
        sheet = book["NODE_INFO"]
        headers = [cell.value for cell in sheet[1]]
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, headers.index("NODE_NAME") + 1).value == "Langfang":
                for field in ("DB_TRACKS", "UB_TRACKS"):
                    sheet.cell(row, headers.index(field) + 1, "[1,2,3,4,5,6,7,8,9,10]/11")
        content = BytesIO()
        book.save(content)
        book.close()
        create_base(client, ids[2], content.getvalue())
        for change in [
            lambda s: setattr(s["B2"], "value", 999),
            lambda s: setattr(s["L3"], "value", 999),
            lambda s: setattr(s["M2"], "value", 0),
            lambda s: setattr(s["J2"], "value", "25:99:00"),
        ]:
            job = preview(client, ids[-1], result_book(change=change)).json()
            assert not job["report"]["can_commit"], job["report"]["issues"]
            assert commit(client, job).status_code == 409

        def swapped_seq(sheet):
            sheet["M2"], sheet["M3"] = sheet["M3"].value, sheet["M2"].value

        job = preview(client, ids[-1], result_book(change=swapped_seq)).json()
        assert any(
            "SEQ" in i["message"] for i in job["report"]["issues"] if i["severity"] == "error"
        )
        _, run = imported(client, ids[-1])
        other = branch(client, "other")
        plan = client.get(f"/api/v1/catalog/{ids[-1]}/deletion-preview").json()
        response = client.request(
            "DELETE",
            f"/api/v1/catalog/{ids[-1]}",
            json={
                "confirmed": True,
                "preview_token": plan["preview_token"],
                "action": "move",
                "target_batch_id": other[-1],
            },
        )
        assert response.status_code == 422, response.text
        assert client.get(f"/api/v1/runs/{run['id']}").json()["owner_category_id"] == ids[-1]


def test_pending_number_race_and_startup_recovery_do_not_overwrite_runs(tmp_path):
    settings = Settings(database=tmp_path / "db")
    app = create_app(settings)
    with TestClient(app) as client:
        ids = branch(client, "race")
        first, second = preview(client, ids[-1]).json(), preview(client, ids[-1]).json()
        assert first["report"]["options"]["run_no"] == second["report"]["options"]["run_no"] == 1
        assert commit(client, first).status_code == 200
        assert commit(client, second).status_code == 409
        pending = preview(client, ids[-1]).json()
        assert pending["report"]["options"]["run_no"] == 2
        with app.state.database.transaction() as c:
            c.execute("UPDATE import_jobs SET state='prepared' WHERE id=?", [pending["job_id"]])
        (tmp_path / "source_files" / f"{pending['job_id']}.xlsx").write_bytes(b"interrupted")
    with TestClient(create_app(settings)) as client:
        assert (
            client.get(f"/api/v1/runs/import/jobs/{pending['job_id']}").json()["state"] == "failed"
        )
        assert not (tmp_path / "source_files" / f"{pending['job_id']}.xlsx").exists()
        assert len(client.get("/api/v1/runs").json()) == 1
