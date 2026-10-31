from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from app.base_timetables.exporting import sheet_parts
from app.core.config import Settings
from app.ingest.base_parser import parse_base
from app.main import create_app
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from sample_workbook import SAMPLE_WORKBOOK
from test_base_import import create_base
from test_base_replace import seed_result
from test_catalog_deletion import branch


@pytest.fixture
def context(tmp_path):
    sample = SAMPLE_WORKBOOK.read_bytes()
    app = create_app(Settings(database=tmp_path / "db"))
    with TestClient(app) as client:
        ids = branch(client, "date-export")
        _, base = create_base(client, ids[3], sample)
        seed_result(app, client, ids[-1], base)
        yield app, client, ids, base, sample


def save(client, url, date_base=None, changes=None, **flags):
    data = client.get(url + "/tables").json()
    payload = {"revision": data["revision"], "changes": changes or [], "date_base": date_base}
    report = client.post(url + "/events/validate", json=payload)
    assert report.status_code == 200, report.text
    payload.update(
        validation_token=report.json()["validation_token"],
        warnings_confirmed=True,
        train_dates_confirmed=True,
        date_base_confirmed=True,
    )
    payload.update(flags)
    return data, report.json(), client.post(url + "/events/edit", json=payload)


def exported(client, url, selected=None, template=None):
    options = client.get(url + "/export-options")
    assert options.status_code == 200, options.text
    options = options.json()
    response = client.post(
        url + "/export",
        params={
            "revision": options["revision"],
            "sheets": ",".join(options["defaults"] if selected is None else selected),
            "replace": template is not None,
        },
        content=template or b"",
    )
    return options, response


@pytest.mark.parametrize("kind", ["base", "run"])
def test_date_only_save_updates_all_dates_atomically(context, kind):
    app, client, _, base, sample = context
    url = f"/api/v1/base-timetables/{base['id']}" if kind == "base" else "/api/v1/runs/result"
    before = client.get(url + "/tables").json()
    untouched = client.get(
        "/api/v1/runs/result/tables"
        if kind == "base"
        else f"/api/v1/base-timetables/{base['id']}/tables"
    ).json()
    _, report, rejected = save(client, url, "2025-06-01", date_base_confirmed=False)
    assert report["can_save"] and report["date_base_change"]["after"] == "2025-06-01"
    assert rejected.status_code == 422
    assert client.get(url + "/tables").json() == before
    _, _, response = save(client, url, "2025-06-01")
    assert response.status_code == 200, response.text
    after = response.json()
    assert after["object"]["date_base"] == "2025-06-01"
    for old, new in zip(before["events"], after["events"], strict=True):
        assert old["id"] == new["id"]
        for field in ("arrival", "departure"):
            assert old[field] == new[field]
            if old[field]:
                assert datetime.fromisoformat(new[field + "_datetime"]) - datetime.fromisoformat(
                    old[field + "_datetime"]
                ) == timedelta(days=7)
    assert (
        client.get(
            "/api/v1/runs/result/tables"
            if kind == "base"
            else f"/api/v1/base-timetables/{base['id']}/tables"
        ).json()["events"]
        == untouched["events"]
    )
    assert Path(base["source_files"][0]["managed_path"]).read_bytes() == sample
    options, response = exported(client, url, ["TTT"])
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content))
    sheet = book[options["sheets"][0]["name"]]
    values = [cell.value for row in sheet for cell in row if isinstance(cell.value, datetime)]
    assert min(d.date() for d in values).isoformat() == "2025-06-01"
    for header in sheet[1]:
        if isinstance(header.value, str) and header.value.endswith(("_A", "_D")):
            assert all(
                row[0].number_format == "hh:mm:ss"
                for row in sheet.iter_rows(min_row=2, min_col=header.column, max_col=header.column)
            )
    book.close()


def test_date_and_cross_day_edit_token_and_invalid_dates(context):
    _, client, _, base, _ = context
    url = f"/api/v1/base-timetables/{base['id']}"
    before = client.get(url + "/tables").json()
    event = next(e for e in before["events"] if e["departure"] is not None)
    report = client.post(
        url + "/events/validate", json={"revision": before["revision"], "date_base": "2025-06-01"}
    ).json()
    assert (
        client.post(
            url + "/events/edit",
            json={
                "revision": before["revision"],
                "date_base": "2025-06-02",
                "validation_token": report["validation_token"],
                "date_base_confirmed": True,
                "train_dates_confirmed": True,
            },
        ).status_code
        == 409
    )
    assert client.get(url + "/tables").json() == before
    changes = [{"event_id": event["id"], "field": "departure", "value": "2025-06-02 01:10:00"}]
    _, report, response = save(client, url, "2025-06-01", changes)
    assert response.status_code == 200, response.text
    saved = next(e for e in response.json()["events"] if e["id"] == event["id"])
    assert saved["departure_datetime"] == "2025-06-02T01:10:00"
    assert saved["departure"]["seconds"] == 90600
    _, exported_file = exported(client, url, ["TTT"])
    assert exported_file.status_code == 200, exported_file.text
    exported_book = load_workbook(BytesIO(exported_file.content))
    source = event["source"]
    exported_cell = exported_book[source["sheet"]].cell(source["row"], source["first_column"] + 1)
    assert exported_cell.value == datetime(2025, 6, 2, 1, 10)
    assert exported_cell.number_format == "hh:mm:ss"
    exported_book.close()
    after = client.get(url + "/tables").json()
    assert (
        client.post(
            url + "/events/edit",
            json={
                "revision": before["revision"],
                "changes": changes,
                "date_base": "2025-06-03",
                "validation_token": report["validation_token"],
                "date_base_confirmed": True,
                "train_dates_confirmed": True,
            },
        ).status_code
        == 409
    )
    assert (
        client.post(
            url + "/events/validate",
            json={"revision": after["revision"], "date_base": "2010-01-01"},
        ).status_code
        == 422
    )
    invalid = client.post(
        url + "/events/validate", json={"revision": after["revision"], "date_base": "9999-12-31"}
    ).json()
    assert not invalid["can_save"]
    assert client.get(url + "/tables").json() == after


def test_defaults_four_tables_preserve_entire_workbook_and_roundtrip(context):
    _, client, _, base, sample = context
    base_url = f"/api/v1/base-timetables/{base['id']}"
    options, response = exported(client, base_url)
    assert response.status_code == 200, response.text
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    book = load_workbook(BytesIO(response.content))
    original = load_workbook(BytesIO(sample))
    assert book.sheetnames == original.sheetnames
    original.close()
    book.close()
    report = parse_base(
        response.content, datetime(2025, 5, 25).date(), "preserve", track_lists=True
    )
    assert not [i for i in report["issues"] if i["severity"] == "error"]
    assert len(report["events"]) == 148 and len(report["nodes"]) == 53
    options, result = exported(client, "/api/v1/runs/result")
    assert options["defaults"] == ["TTT", "LINE_INFO", "NODE_INFO", "MINIMUM_RUN_TIME"]
    assert result.status_code == 200, result.text
    book = load_workbook(BytesIO(result.content))
    assert book.sheetnames == ["TTT1base", "LINE_INFO", "NODE_INFO", "MINIMUM_RUN_TIME"]
    book.close()
    _, result = exported(client, "/api/v1/runs/result", ["TTT", "NODE_INFO"])
    book = load_workbook(BytesIO(result.content))
    assert book.sheetnames == ["TTT1base", "LINE_INFO", "NODE_INFO", "MINIMUM_RUN_TIME"]
    book.close()
    assert Path(base["source_files"][0]["managed_path"]).read_bytes() == sample


@pytest.mark.parametrize("epoch1904", [False, True])
def test_overwrite_only_selected_parts_and_keep_epoch(context, epoch1904):
    _, client, _, base, sample = context
    source = load_workbook(BytesIO(sample))
    extra = source.create_sheet("用户备注")
    extra["A1"] = "应保留"
    extra["B2"] = "=1+2"
    extra.column_dimensions["A"].width = 37
    if epoch1904:
        source.epoch = CALENDAR_MAC_1904
    buffer = BytesIO()
    source.save(buffer)
    source.close()
    original = buffer.getvalue()
    url = f"/api/v1/base-timetables/{base['id']}"
    _, _, saved = save(client, url, "2025-06-01")
    assert saved.status_code == 200
    _, response = exported(client, url, ["TTT"], original)
    assert response.status_code == 200, response.text
    with ZipFile(BytesIO(original)) as before, ZipFile(BytesIO(response.content)) as after:
        selected_part = sheet_parts(before)["TTT1base"]
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name not in {selected_part, "xl/styles.xml"}:
                assert before.read(name) == after.read(name), name
    merged = load_workbook(BytesIO(response.content))
    assert merged["用户备注"]["B2"].value == "=1+2"
    assert merged["用户备注"].column_dimensions["A"].width == 37
    dates = [c.value for row in merged["TTT1base"] for c in row if isinstance(c.value, datetime)]
    assert min(t.date() for t in dates).isoformat() == "2025-06-01"
    for header in merged["TTT1base"][1]:
        if isinstance(header.value, str) and header.value.endswith(("_A", "_D")):
            assert all(
                row[0].number_format == "hh:mm:ss"
                for row in merged["TTT1base"].iter_rows(
                    min_row=2, min_col=header.column, max_col=header.column
                )
            )
    if epoch1904:
        assert merged.epoch == CALENDAR_MAC_1904
    merged.close()


def test_invalid_selection_missing_source_and_revision_conflicts(context):
    _, client, _, base, sample = context
    url = f"/api/v1/base-timetables/{base['id']}"
    options = client.get(url + "/export-options").json()
    for sheets in ("", "OTHER", "TTT,TTT"):
        response = client.post(
            url + "/export", params={"revision": options["revision"], "sheets": sheets}
        )
        assert response.status_code == 422
    assert (
        client.post(url + "/export", params={"revision": "stale", "sheets": "TTT"}).status_code
        == 409
    )
    assert (
        client.post(
            url + "/export",
            params={"revision": options["revision"], "sheets": "TTT", "replace": True},
        ).status_code
        == 422
    )
    assert (
        client.post(
            url + "/export",
            params={"revision": options["revision"], "sheets": "TTT", "replace": True},
            content=b"bad",
        ).status_code
        == 422
    )
    source = load_workbook(BytesIO(sample))
    del source["TTT1base"]
    buffer = BytesIO()
    source.save(buffer)
    source.close()
    _, response = exported(client, url, ["TTT"], buffer.getvalue())
    assert response.status_code == 422
    before = client.get(url + "/tables").json()
    _, result = exported(client, url, ["NODE_INFO"])
    assert result.status_code == 200
    assert client.get(url + "/tables").json() == before


def test_date_save_database_failure_rolls_back_events_and_metadata(context, monkeypatch):
    from contextlib import contextmanager

    app, client, _, _, _ = context
    url = "/api/v1/runs/result"
    before = client.get(url + "/tables").json()
    payload = {"revision": before["revision"], "date_base": "2025-06-01", "changes": []}
    report = client.post(url + "/events/validate", json=payload).json()
    payload.update(
        validation_token=report["validation_token"],
        warnings_confirmed=True,
        train_dates_confirmed=True,
        date_base_confirmed=True,
    )
    database = app.state.database
    original = database.transaction

    class Broken:
        def __init__(self, c):
            self.c = c

        def execute(self, sql, *args):
            if sql.startswith("UPDATE run_working_data"):
                raise RuntimeError("date save failure")
            return self.c.execute(sql, *args)

    @contextmanager
    def broken():
        with original() as c:
            yield Broken(c)

    monkeypatch.setattr(database, "transaction", broken)
    with pytest.raises(RuntimeError, match="date save failure"):
        client.post(url + "/events/edit", json=payload)
    monkeypatch.setattr(database, "transaction", original)
    assert client.get(url + "/tables").json() == before
