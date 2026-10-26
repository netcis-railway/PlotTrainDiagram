from copy import deepcopy
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from app.base_timetables.ordering import route_nodes
from app.core.config import Settings
from app.main import create_app
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sample_workbook import SAMPLE_WORKBOOK
from test_base_import import create_base
from test_base_replace import seed_result
from test_catalog_deletion import branch


def reordered_sample():
    original = SAMPLE_WORKBOOK.read_bytes()
    book = load_workbook(BytesIO(original))
    sheet = book["TTT1base"]
    rows = list(sheet.values)
    header = rows[0]
    groups = list(range(5, len(header), 4))[::-1]

    def reordered(row):
        return list(row[:5]) + [v for i in groups for v in row[i : i + 4]]

    sheet.delete_rows(1, sheet.max_row)
    sheet.append(reordered(header))
    station_names = []
    for row in reversed(rows[1:]):
        sheet.append(reordered(row))
        if isinstance(row[4], str):
            station_names.append(row[4])
    courses = [str(header[i])[:-2] for i in groups]
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue(), station_names, courses


def test_base_order_uses_excel_rows_and_columns_not_database_or_numeric_sort(tmp_path):
    sample, names, courses = reordered_sample()
    app = create_app(Settings(database=tmp_path / "db"))
    with TestClient(app) as client:
        ids = branch(client, "order")
        _, base = create_base(client, ids[3], sample)
        url = f"/api/v1/base-timetables/{base['id']}/tables"
        result = client.get(url).json()
        node_names = {n["id"]: n["node_name"] for n in result["nodes"]}
        train_names = {t["id"]: t["train_course_id"] for t in result["trains"]}
        order = result["original_order"]
        assert order["source"] == "excel"
        assert [node_names[n] for n in order["nodes"]][: len(names)] == names
        assert [train_names[t] for t in order["trains"]] == courses
        assert all(n["node_ind"] and n["line_inds"] for n in result["nodes"])
        assert result["display_lines"]
        assert all(
            line["line_ind"] > 0 and line["line_name_chn"] for line in result["display_lines"]
        )
        expected_members = {}
        for member in result["infrastructure"]["members"]:
            expected_members.setdefault(member["NODE_NAME"], set()).add(member["LINE_IND"])
        for node in result["nodes"]:
            assert expected_members[node["node_name"]] <= set(node["line_inds"])
        assert client.get(url).json() == result


def test_run_metadata_and_original_order_with_own_excel(tmp_path):
    sample, names, courses = reordered_sample()
    app = create_app(Settings(database=tmp_path / "db"))
    with TestClient(app) as client:
        ids = branch(client, "order")
        _, base = create_base(client, ids[3], sample)
        seed_result(app, client, ids[-1], base)
        source = tmp_path / "source_files" / "result.xlsx"
        source.write_bytes(sample)
        with app.state.database.transaction() as c:
            c.execute(
                "INSERT INTO source_files VALUES (?,?,?,?,?,?)",
                [str(uuid4()), "result", str(uuid4()), "result.xlsx", "test", source.name],
            )
        result = client.get("/api/v1/runs/result/tables").json()
        stations = {n["id"]: n for n in result["nodes"]}
        trains = {t["id"]: t for t in result["trains"]}
        order = result["original_order"]
        assert order["source"] == "excel"
        expected = [n for n in names if n in {s["node_name"] for s in stations.values()}]
        assert [stations[n]["node_name"] for n in order["nodes"]] == expected
        assert [trains[t]["train_course_id"] for t in order["trains"]] == courses
        assert all(n["node_ind"] and n["line_inds"] for n in stations.values())
        # Ordinary runs use the nearest effective base's actual line names.
        base_tables = client.get(f"/api/v1/base-timetables/{base['id']}/tables").json()
        assert result["display_lines"] == base_tables["display_lines"]
        assert len(result["display_lines"]) == 12
        route = {n["node_name"]: n for n in result["route_nodes"]}
        assert len(route) == len(base_tables["nodes"])
        assert len(route) > len(stations)
        for station in stations.values():
            assert route[station["node_name"]] == {
                k: station[k]
                for k in ("id", "node_name", "node_name_chn", "node_ind", "line_inds", "line_seqs")
            }


def test_missing_original_file_falls_back_without_changing_revision_or_saved_data(tmp_path):
    sample, _, _ = reordered_sample()
    app = create_app(Settings(database=tmp_path / "db"))
    with TestClient(app) as client:
        ids = branch(client, "order")
        _, base = create_base(client, ids[3], sample)
        url = f"/api/v1/base-timetables/{base['id']}/tables"
        before = client.get(url).json()
        Path(base["source_files"][0]["managed_path"]).unlink()
        after = client.get(url).json()
        assert after["revision"] == before["revision"]
        assert after["events"] == before["events"]
        assert after["original_order"]["source"] == "provenance"
        assert after["original_order"]["notice"]
        assert len(after["original_order"]["trains"]) == len(before["trains"])


def test_route_metadata_retains_unvisited_positions_and_maps_own_ids_without_mutation():
    infrastructure = {
        "nodes": [
            {"id": "base-a", "node_name": "a", "node_ind": 1},
            {"id": "base-b", "node_name": "b", "node_ind": 100},
        ],
        "members": [
            {"NODE_NAME": "a", "LINE_IND": 1, "NODE_SEQ_OF_LINE": 1},
            {"NODE_NAME": "b", "LINE_IND": 1, "NODE_SEQ_OF_LINE": 2},
            {"NODE_NAME": "b", "LINE_IND": 2, "NODE_SEQ_OF_LINE": 5},
        ],
    }
    own = [{"id": "run-a", "node_name": "a", "node_name_chn": "A", "node_ind": 1, "line_inds": [1]}]
    original = deepcopy((own, infrastructure))
    route = {n["node_name"]: n for n in route_nodes(own, infrastructure)}
    assert route["a"]["id"] == "run-a"
    assert route["b"]["id"] == "base-b"
    assert route["b"]["node_ind"] == 100
    assert route["b"]["line_inds"] == [1, 2]
    assert route["b"]["line_seqs"] == {"1": 2, "2": 5}
    assert (own, infrastructure) == original
