from collections import defaultdict
from datetime import date
from io import BytesIO

import pytest
from app.core.config import Settings
from app.ingest.base_parser import parse_base
from app.ingest.sequences import sequence_problem
from app.ingest.tracks import check_event_track
from app.main import create_app
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from test_base_import import create_base
from test_catalog_deletion import branch
from test_run_import import commit as commit_run
from test_run_import import preview as preview_run
from test_track_usage import add_station_yard, content, track_book


def shuffled_offset_book():
    book = track_book()
    sheet = book["TTT1base"]
    for row in range(2, sheet.max_row + 1):
        for col in range(9, sheet.max_column + 1, 4):
            if sheet.cell(row, col).value is not None:
                sheet.cell(row, col, sheet.cell(row, col).value + 37)
                sheet.cell(row, col - 1, [16, 9, 11, 10][row % 4])
    rows = list(sheet.values)[1:]
    for index, values in enumerate(reversed(rows), 2):
        for col, value in enumerate(values, 1):
            sheet.cell(index, col).value = value
    return book


@pytest.mark.parametrize("values", [[1], [37], [2147483647], [7, 5, 6], [1003, 1001, 1002]])
def test_train_sequence_accepts_any_positive_start_and_source_order(values):
    assert sequence_problem(values) is None


@pytest.mark.parametrize(
    "values,word",
    [([5, 5], "重复"), ([5, 7], "缺少6"), ([100, 104], "缺少101至103"), ([0], "正整数")],
)
def test_sequence_duplicates_gaps_and_nonpositive_values_are_rejected(values, word):
    assert word in sequence_problem(values)[0]


def test_track_membership_is_station_or_not_all_lines_and_never_other_station():
    members = [
        {"node_name": "S", "line_ind": 1, "DB_TRACKS": "[6,7]/5", "UB_TRACKS": "[1,2,3]/4"},
        {"node_name": "S", "line_ind": 3, "DB_TRACKS": "[12,13]/11", "UB_TRACKS": "[8,9]/10"},
        {"node_name": "OTHER", "line_ind": 1, "DB_TRACKS": "[99]", "UB_TRACKS": "[99]"},
    ]
    for track in (1, 5, 6, 9, 11, 13):
        assert check_event_track(members, "S", track)[0] == "valid"
    state, numbers, lines = check_event_track(members, "S", 99)
    assert state == "invalid" and numbers == list(range(1, 14)) and lines == [1, 3]
    members[1]["UB_TRACKS"] = None
    assert check_event_track(members, "S", 5)[0] == "valid"
    assert check_event_track(members, "S", 99)[0] == "unknown"


def test_import_keeps_offset_sequences_in_scrambled_rows_and_ignores_track_order():
    book = shuffled_offset_book()
    report = parse_base(content(book), date(2025, 5, 25), "rebase", track_lists=True)
    book.close()
    assert not [i for i in report["issues"] if i["severity"] == "error"]
    assert len(report["events"]) == 148
    grouped = defaultdict(list)
    for event in report["events"]:
        grouped[event["train_course_id"]].append(event["visit_seq"])
    assert all(min(values) == 38 and not sequence_problem(values) for values in grouped.values())
    train = next(t for t in report["trains"] if t["train_course_id"] == "G173")
    assert train["adopted_seq"] == 38 and train["source"]["row"] != 2


def test_bad_track_does_not_create_sequence_gap_and_errors_keep_exact_cells():
    book = track_book()
    sheet = book["TTT1base"]
    sheet["L3"] = 999
    report = parse_base(content(book), date(2025, 5, 25), "rebase", track_lists=True)
    errors = [i for i in report["issues"] if i["severity"] == "error"]
    assert len(errors) == 1
    assert (
        errors[0]["cell"] == "L3" and errors[0]["field"] == "G173_G" and errors[0]["value"] == 999
    )
    sheet["M3"] = 100
    report = parse_base(content(book), date(2025, 5, 25), "rebase", track_lists=True)
    errors = [i for i in report["issues"] if i["severity"] == "error"]
    assert len(errors) == 2
    gap = next(i for i in errors if i["field"] == "G173_SEQ")
    assert gap["cell"] == "M4" and gap["value"] == 3 and "缺少2" in gap["message"]
    sheet["L3"] = 10
    sheet["M3"] = 1
    report = parse_base(content(book), date(2025, 5, 25), "rebase", track_lists=True)
    errors = [i for i in report["issues"] if i["severity"] == "error"]
    assert len(errors) == 1 and errors[0]["field"] == "G173_SEQ" and errors[0]["cell"] == "M3"
    assert "重复" in errors[0]["message"]
    book.close()


def test_base_run_import_edit_export_and_plot_preserve_nonone_seq_and_station_or(tmp_path):
    book = shuffled_offset_book()
    # A different line owns track 99 at the same station, even if TTT lists only line 1.
    add_station_yard(book)
    sheet = book["TTT1base"]
    row = next(r for r in range(2, sheet.max_row + 1) if sheet.cell(r, 5).value == "TianjinNan")
    sheet.cell(row, 4, 1)
    sheet.cell(row, 12, 99)
    source = content(book)
    book.close()
    app = create_app(Settings(database=tmp_path / "db"))
    with TestClient(app) as client:
        ids = branch(client, "seq-or")
        _, base = create_base(client, ids[2], source)
        job = preview_run(client, ids[-1], source).json()
        assert job["report"]["can_commit"], job["report"]["issues"]
        response = commit_run(client, job)
        assert response.status_code == 200, response.text
        run = response.json()
        for kind, object_id, url in (
            ("base", base["id"], f"/api/v1/base-timetables/{base['id']}"),
            ("run", run["id"], f"/api/v1/runs/{run['id']}"),
        ):
            tables = client.get(url + "/tables").json()
            seqs = {e["id"]: e["visit_seq"] for e in tables["events"]}
            assert min(seqs.values()) == 38
            event = next(
                e
                for e in tables["events"]
                if e["node_name"] == "TianjinNan" and e["train_course_id"] == "G173"
            )
            body = {
                "revision": tables["revision"],
                "changes": [{"event_id": event["id"], "field": "track_code", "value": 99}],
            }
            valid = client.post(url + "/events/validate", json=body).json()
            assert valid["can_save"], valid["issues"]
            saved = client.post(
                url + "/events/edit",
                json=body
                | {
                    "validation_token": valid["validation_token"],
                    "warnings_confirmed": True,
                    "train_dates_confirmed": True,
                },
            )
            assert saved.status_code == 200, saved.text
            latest = client.get(url + "/tables").json()
            assert {e["id"]: e["visit_seq"] for e in latest["events"]} == seqs
            export = client.post(
                url + "/export", params={"revision": latest["revision"], "sheets": "TTT"}
            )
            assert export.status_code == 200, export.text
            out = load_workbook(BytesIO(export.content))
            rows = list(out.active.values)
            values = [
                row[col]
                for row in rows[1:]
                for col in range(8, len(rows[0]), 4)
                if row[col] is not None
            ]
            out.close()
            assert min(values) == 38
            plot = client.post(
                "/api/v1/diagram/data",
                json={
                    "object_type": kind,
                    "object_id": object_id,
                    "revision": latest["revision"],
                    "line_ind": 1,
                    "overlay_base": kind == "run",
                },
            ).json()
            assert plot["can_plot"] and not any(
                i["code"] in ("sequence", "sequence_gap") for i in plot["issues"]
            )
            resource = next(r for r in plot["track_stations"] if r["node_id"] == event["node_id"])
            assert 99 not in [t["number"] for t in resource["tracks"]]  # plot stays line-scoped
            # Missing one visit number is still blocked in the editor; no implicit renumbering.
            body = {
                "revision": latest["revision"],
                "changes": [{"event_id": event["id"], "field": "visit_seq", "value": 1000}],
            }
            bad = client.post(url + "/events/validate", json=body).json()
            assert not bad["can_save"] and any(
                i["field"] == "visit_seq" and "不连续" in i["message"] for i in bad["issues"]
            )
