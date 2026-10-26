from datetime import date
from io import BytesIO

import pytest
from app.core.config import Settings
from app.ingest.base_parser import parse_base
from app.ingest.inspection import inspect_workbook
from app.main import create_app
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from test_base_import import commit, create_base, preview
from test_catalog_deletion import branch
from test_run_import import commit as run_commit
from test_run_import import preview as run_preview
from test_track_repair import workbook


def source(missing=("G1",), both=False, result=False, gaps=True):
    book = load_workbook(BytesIO(workbook(result=result)))
    table = book["TTT1"]
    if gaps:
        table.insert_rows(3, 2)
    for course in reversed(("G1", "G2/G3")):
        if course not in missing:
            continue
        headers = [cell.value for cell in table[1]]
        table.delete_cols(headers.index(course + "_SEQ") + 1)
        if both:
            headers = [cell.value for cell in table[1]]
            table.delete_cols(headers.index(course + "_G") + 1)
    out = BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


@pytest.mark.parametrize(
    "missing,both", [(("G1",), False), (("G1", "G2/G3"), False), (("G1", "G2/G3"), True)]
)
def test_structural_preflight_separates_each_missing_seq_and_g(missing, both):
    data = source(missing, both)
    sheet = inspect_workbook(data)["sheets"][0]
    assert sheet["trains"] == ["G1", "G2/G3"]
    report = parse_base(data, date(2025, 5, 25), "rebase", track_lists=True)
    seq = [i for i in report["issues"] if i.get("field", "").endswith("_SEQ")]
    assert {i["field"] for i in seq} == {course + "_SEQ" for course in missing}
    assert all(i["row"] == 1 and "缺少序号列" in i["message"] for i in seq)
    assert len([i for i in sheet["issues"] if i.get("field", "").endswith("_G")]) == (
        len(missing) if both else 0
    )


@pytest.mark.parametrize("kind", ["base", "run"])
def test_repair_numbers_only_timed_rows_preserves_existing_seq_and_original_then_repreflights(
    tmp_path, kind
):
    app = create_app(Settings(database=tmp_path / "seq.duckdb"))
    data = source(result=kind == "run")
    with TestClient(app) as client:
        ids = branch(client, "seq-repair")
        if kind == "run":
            # The existing base uses the same generated sequence for G1 and retains G2's 41/42.
            foundation = load_workbook(BytesIO(workbook()))
            foundation["TTT1"]["I2"] = 1
            foundation["TTT1"]["I3"] = 2
            out = BytesIO()
            foundation.save(out)
            foundation.close()
            create_base(client, ids[2], out.getvalue())
        create = run_preview if kind == "run" else preview
        job = create(client, ids[-1] if kind == "run" else ids[2], data).json()
        assert not job["report"]["can_commit"]
        url = f"/api/v1/ingest/jobs/{job['job_id']}/sequence-repair"
        response = client.get(url + "/context")
        assert response.status_code == 200, response.text
        context = response.json()
        missing = next(t for t in context["trains"] if t["missing_seq"])
        existing = next(t for t in context["trains"] if not t["missing_seq"])
        assert missing["visits"] == 2 and missing["first_row"] == 2 and missing["last_row"] == 5
        body = {"repair_token": context["repair_token"], "train_ids": [missing["id"]]}
        assert client.post(url + "/apply", json=body | {"repair_token": "bad"}).status_code == 409
        assert (
            client.post(url + "/apply", json=body | {"train_ids": [existing["id"]]}).status_code
            == 422
        )
        assert (
            client.post(
                url + "/apply", json=body | {"train_ids": [missing["id"], missing["id"]]}
            ).status_code
            == 422
        )
        repaired = client.post(url + "/apply", json=body)
        assert (
            repaired.status_code == 200 and "spreadsheetml" in repaired.headers["content-type"]
        ), repaired.text
        book = load_workbook(BytesIO(repaired.content))
        table = book["TTT1"]
        headers = [c.value for c in table[1]]
        first, second = headers.index("G1_SEQ") + 1, headers.index("G2/G3_SEQ") + 1
        assert [table.cell(row, first).value for row in range(2, 6)] == [1, None, None, 2]
        assert [table.cell(row, second).value for row in range(2, 6)] == [41, None, None, 42]
        original = load_workbook(BytesIO(data))
        for field in [c.value for c in original["TTT1"][1]]:
            old_col = [c.value for c in original["TTT1"][1]].index(field) + 1
            new_col = headers.index(field) + 1
            assert [original["TTT1"].cell(row, old_col).value for row in range(1, 6)] == [
                table.cell(row, new_col).value for row in range(1, 6)
            ]
        original.close()
        book.close()
        assert client.get(url + "/context").json()["repair_token"] == context["repair_token"]
        job2 = create(client, ids[-1] if kind == "run" else ids[2], repaired.content).json()
        assert job2["report"]["can_commit"], job2["report"]["issues"]
        saved = (run_commit if kind == "run" else commit)(client, job2)
        assert saved.status_code == 200, saved.text
        tables = client.get(
            f"/api/v1/{'runs' if kind == 'run' else 'base-timetables'}/{saved.json()['id']}/tables"
        ).json()
        assert sorted(e["visit_seq"] for e in tables["events"]) == [1, 2, 41, 42]
        assert (
            client.get(f"/api/v1/ingest/jobs/{job2['job_id']}/sequence-repair/context").status_code
            == 422
        )


def test_seq_then_track_repairs_handle_both_missing_without_column_shift(tmp_path):
    app = create_app(Settings(database=tmp_path / "both.duckdb"))
    with TestClient(app) as client:
        ids = branch(client, "both-repairs")
        data = source(missing=("G1", "G2/G3"), both=True)
        job = preview(client, ids[2], data).json()
        root = f"/api/v1/ingest/jobs/{job['job_id']}"
        track = client.get(root + "/track-repair/context")
        assert track.status_code == 422 and "先点击修复SEQ" in track.json()["detail"]
        context = client.get(root + "/sequence-repair/context").json()
        fixed = client.post(
            root + "/sequence-repair/apply",
            json={
                "repair_token": context["repair_token"],
                "train_ids": [t["id"] for t in context["trains"]],
            },
        )
        assert fixed.status_code == 200, fixed.text
        job2 = preview(client, ids[2], fixed.content).json()
        assert not job2["report"]["can_commit"]
        assert {i["field"] for i in job2["report"]["issues"] if i["field"].endswith("_G")} == {
            "G1_G",
            "G2/G3_G",
        }
        root = f"/api/v1/ingest/jobs/{job2['job_id']}"
        ctx = client.get(root + "/track-repair/context").json()
        assert "repair_token" in ctx, ctx
        fixed2 = client.post(
            root + "/track-repair/apply",
            json={
                "repair_token": ctx["repair_token"],
                "selections": [
                    {
                        "train_id": t["id"],
                        "lines": [{"line_ind": 1, "directions": [t["direction"]]}],
                    }
                    for t in ctx["trains"]
                ],
            },
        )
        job3 = preview(client, ids[2], fixed2.content).json()
        assert job3["report"]["can_commit"], job3["report"]["issues"]
        assert [e["visit_seq"] for e in job3["report"]["events"]] == [1, 1, 2, 2]


def test_time_presence_includes_one_sided_and_zero_but_skips_blank_and_track_only_rows():
    from app.ingest.sequence_repair import sequence_rows
    from openpyxl import Workbook

    book = Workbook()
    table = book.active
    table.append(["A", "D", "G"])
    table.append(["07:00:00", None, None])
    table.append([None, "07:01:00", None])
    table.append([None, None, 99])
    table.append([" ", None, None])
    table.append([0, None, None])
    assert sequence_rows(table, 1) == [2, 3, 6]
    book.close()


def test_repaired_column_ignores_track_only_row_without_hiding_real_seq_or_formula_errors():
    book = load_workbook(BytesIO(workbook()))
    table = book["TTT1"]
    table.insert_rows(3)
    for column, value in enumerate(["丙", 3, "R", 1, "R"], 1):
        table.cell(3, column, value)
    table.cell(3, 8, 6)  # G-only, with no A/D or SEQ: not a visit.
    data = BytesIO()
    book.save(data)
    report = parse_base(data.getvalue(), date(2025, 5, 25), "rebase", track_lists=True)
    assert not [i for i in report["issues"] if i["severity"] == "error"], report["issues"]
    assert len(report["events"]) == 4
    assert any(
        i["field"] == "G1_G" and i["row"] == 3 and "不作为列车访问" in i["message"]
        for i in report["issues"]
    )
    table.cell(3, 9, 43)  # Explicit SEQ still retains a visit, even with missing times.
    data = BytesIO()
    book.save(data)
    report = parse_base(data.getvalue(), date(2025, 5, 25), "rebase", track_lists=True)
    assert len(report["events"]) == 5
    table.cell(3, 9).value = None
    table.cell(3, 8, "=6")
    data = BytesIO()
    book.save(data)
    book.close()
    report = parse_base(data.getvalue(), date(2025, 5, 25), "rebase", track_lists=True)
    assert any(
        i["field"] == "G1_G" and i["severity"] == "error" and "公式" in i["message"]
        for i in report["issues"]
    )
