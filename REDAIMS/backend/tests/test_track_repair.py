from copy import deepcopy
from datetime import date, datetime
from io import BytesIO

import pytest
from app.base_timetables.track_repair import (
    TrainChoice,
    allocate,
    default_direction,
    interval,
    overlaps,
    repair_context,
)
from app.core.config import Settings
from app.ingest.base_parser import parse_base
from app.ingest.inspection import inspect_workbook
from app.main import create_app
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from test_base_import import commit, create_base, preview
from test_catalog_deletion import branch
from test_run_import import commit as run_commit
from test_run_import import preview as run_preview


def workbook(missing=False, result=False):
    book = Workbook()
    table = book.active
    table.title = "TTT1"
    names = ["G1", "G2/G3"]
    table.append(
        ["NODE_NAME_CHN", "NODE_IND", "NODE_CODE", "LINE_IND", "NODE_NAME"]
        + [f"{name}_{suffix}" for name in names for suffix in ("A", "D", "G", "SEQ")]
    )
    table.append(
        [
            "甲",
            1,
            "P",
            1,
            "P",
            None,
            datetime(2025, 5, 25, 7),
            None,
            41,
            None,
            datetime(2025, 5, 25, 7, 0, 30),
            None,
            41,
        ]
    )
    table.append(
        [
            "乙",
            2,
            "Q",
            1,
            "Q",
            datetime(2025, 5, 25, 7, 2),
            datetime(2025, 5, 25, 7, 3),
            None,
            42,
            datetime(2025, 5, 25, 7, 2, 30),
            datetime(2025, 5, 25, 7, 3, 30),
            None,
            42,
        ]
    )
    if missing:
        table.delete_cols(12)
        table.delete_cols(8)
    if not result:
        line = book.create_sheet("LINE_INFO")
        line.append(["LINE_IND", "LINE_CODE", "LINE_NAME_CHN", "NUM_NODE"])
        line.append([1, "L", "测试线路", 2])
        nodes = book.create_sheet("NODE_INFO")
        nodes.append(
            [
                "LINE_IND",
                "LINE_CODE",
                "LINE_NAME_CHN",
                "NODE_SEQ_OF_LINE",
                "NODE_NAME_CHN",
                "NODE_IND",
                "NODE_CODE",
                "NODE_NAME",
                "DB_TRACKS",
                "UB_TRACKS",
            ]
        )
        nodes.append([1, "L", "测试线路", 1, "甲", 1, "P", "P", "[9,10]/11", "[6,7]/8"])
        nodes.append([1, "L", "测试线路", 2, "乙", 2, "Q", "Q", "[9,10]/11", "[6,7]/8"])
        minimum = book.create_sheet("MINIMUM_RUN_TIME")
        minimum.append(
            [
                "LINK_START_NODE_CODE",
                "LINK_END_NODE_CODE",
                "START_ACTIVITY",
                "END_ACTIVITY",
                "MINIMUM_RUN_TIME",
                "MINIMUM_RUN_TIME_SEC",
            ]
        )
        minimum.append(["P", "Q", "PASS", "PASS", "00:02:00", 120])
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def choice(tid, directions=("UB",), line=1):
    return TrainChoice(train_id=tid, lines=[{"line_ind": line, "directions": list(directions)}])


def small_doc():
    def event(tid, start=10, end=20, track=None):
        return {
            "id": tid,
            "train_id": tid,
            "node_name": "P",
            "train_course_id": tid,
            "arrival": {"seconds": start},
            "departure": {"seconds": end},
            "track_code": track,
        }

    doc = {
        "trains": [{"id": "G1", "train_course_id": "G1"}, {"id": "G3", "train_course_id": "G3"}],
        "nodes": [{"node_name": "P", "node_name_chn": "甲"}],
        "events": [event("G1"), event("G3")],
    }
    infra = {
        "lines": [{"LINE_IND": 1}],
        "members": [
            {"node_name": "P", "line_ind": 1, "UB_TRACKS": "[6,7]/8", "DB_TRACKS": "[9]/10"}
        ],
    }
    return doc, infra


@pytest.mark.parametrize(
    "course,direction",
    [("G109", "DB"), ("G108", "UB"), ("G1914/G1911", "UB"), ("D0001", "DB"), ("NoNumber", None)],
)
def test_default_direction_obeys_user_parity_and_first_course(course, direction):
    assert default_direction(course) == direction


def test_allocator_uses_selected_resources_and_nonoverlap_without_mutating_document():
    doc, infra = small_doc()
    before = deepcopy(doc)
    result = allocate(doc, infra, [choice("G1"), choice("G3")])
    assert result["can_apply"] and {c["value"] for c in result["changes"]} == {6, 7}
    assert doc == before
    assert overlaps((10, 20), (19, 30)) and not overlaps((10, 20), (20, 30))
    doc["events"][1]["arrival"]["seconds"] = 20
    doc["events"][1]["departure"]["seconds"] = 30
    result = allocate(doc, infra, [choice("G1"), choice("G3")])
    assert result["can_apply"] and {c["value"] for c in result["changes"]} == {6}


def test_allocator_keeps_unselected_trains_fixed_across_lines_and_prioritizes_roles():
    doc, infra = small_doc()
    doc["events"][1]["track_code"] = 6
    # Reservation follows the selected event in source order.
    result = allocate(doc, infra, [choice("G1")])
    assert result["can_apply"] and result["changes"][0]["value"] == 7
    infra["members"].append(
        {"node_name": "P", "line_ind": 2, "UB_TRACKS": "[6]/6", "DB_TRACKS": "[]"}
    )
    infra["lines"].append({"LINE_IND": 2})
    result = allocate(doc, infra, [choice("G1", line=2)])
    assert result["can_apply"] and result["changes"] == []
    assert any("容量不足" in warning["message"] for warning in result["warnings"])
    doc["events"][0]["departure"]["seconds"] = 10
    result = allocate(doc, infra, [choice("G1")])
    assert result["changes"][0]["value"] == 8  # passage prefers main line
    assert interval(doc["events"][0]) == (10, 11)


def test_allocator_backtracks_and_reports_capacity_time_unknown_and_bad_selection():
    doc, infra = small_doc()
    infra["members"][0].update(UB_TRACKS="[6]", DB_TRACKS="[7]")
    # Flexible train uses two selected lines, each with exactly one direction.
    infra["lines"].append({"LINE_IND": 2})
    infra["members"].append({"node_name": "P", "line_ind": 2, "DB_TRACKS": "[7]"})
    flexible = TrainChoice(
        train_id="G1",
        lines=[{"line_ind": 1, "directions": ["UB"]}, {"line_ind": 2, "directions": ["DB"]}],
    )
    result = allocate(doc, infra, [flexible, choice("G3")])
    assert result["can_apply"]
    assert {c["event_id"]: c["value"] for c in result["changes"]} == {"G1": 7, "G3": 6}
    result = allocate(doc, infra, [choice("G1"), choice("G3")])
    assert result["can_apply"] and len(result["changes"]) == 1
    assert any("容量不足" in warning["message"] for warning in result["warnings"])
    infra["members"][0]["UB_TRACKS"] = None
    skipped = allocate(doc, infra, [choice("G1")])
    assert skipped["can_apply"] and skipped["changes"] == []
    assert any("股道资料未知" in warning["message"] for warning in skipped["warnings"])
    doc["events"][0]["departure"]["seconds"] = 9
    assert "出发早于" in allocate(doc, infra, [choice("G1")])["issues"][0]["message"]
    with pytest.raises(ValueError, match="不存在或重复"):
        allocate(doc, infra, [choice("unknown")])
    with pytest.raises(ValueError, match="线路不存在"):
        allocate(doc, infra, [choice("G1", line=99)])


@pytest.mark.parametrize("split_fields", [False, True])
def test_unknown_selected_direction_warns_and_preserves_its_fixed_occupation(split_fields):
    doc, infra = small_doc()
    doc["events"][0]["track_code"] = 6
    infra["members"][0].update(UB_TRACKS=None, DB_TRACKS="[6,7]")
    if split_fields:
        infra["members"][0].update(
            TRACKS_SCHEMA_VERSION=2,
            UB_STOP_TRACKS=None,
            UB_MAIN_TRACK=None,
            DB_STOP_TRACKS=[6, 7],
            DB_MAIN_TRACK=None,
        )
    before = deepcopy(doc)
    result = allocate(doc, infra, [choice("G1"), choice("G3", ("DB",))])
    assert result["can_apply"] and not result["issues"]
    assert result["changes"] == [{"event_id": "G3", "field": "track_code", "value": 7}]
    assert len(result["warnings"]) == 1
    assert result["warnings"][0]["train_course_id"] == "G1"
    assert "股道资料未知" in result["warnings"][0]["message"]
    assert doc == before


def test_unknown_visits_do_not_block_repair_at_other_stations_or_invent_a_track():
    doc, infra = small_doc()
    doc["nodes"].append({"node_name": "JinanXi", "node_name_chn": "济南西"})
    infra["members"].append(
        {"node_name": "JinanXi", "line_ind": 1, "UB_TRACKS": None, "DB_TRACKS": None}
    )
    for event, course in zip(doc["events"][:], ["G2660/G2661", "G1909/G1908"]):
        event["train_course_id"] = course
        doc["events"].append(event | {"id": event["id"] + "-unknown", "node_name": "JinanXi"})
    before = deepcopy(doc)
    result = allocate(doc, infra, [choice("G1"), choice("G3")])
    assert result["can_apply"] and not result["issues"]
    assert {c["event_id"] for c in result["changes"]} == {"G1", "G3"}
    unknown = [w for w in result["warnings"] if "股道资料未知" in w["message"]]
    assert {w["train_course_id"] for w in unknown} == {"G2660/G2661", "G1909/G1908"}
    assert all(w["node_name"] == "JinanXi" and w["node_name_chn"] == "济南西" for w in unknown)
    assert doc == before


def test_capacity_and_explicit_empty_resources_warn_and_preserve_existing_values():
    doc, infra = small_doc()
    doc["events"][0]["track_code"] = 6
    infra["members"][0].update(UB_TRACKS=None, DB_TRACKS="[6]")
    result = allocate(doc, infra, [choice("G1"), choice("G3", ("DB",))])
    assert result["can_apply"] and result["changes"] == []
    assert any("容量不足" in warning["message"] for warning in result["warnings"])
    assert any("股道资料未知" in warning["message"] for warning in result["warnings"])
    infra["members"][0]["UB_TRACKS"] = "[]"
    result = allocate(doc, infra, [choice("G1")])
    assert result["can_apply"] and result["changes"] == []
    assert any("没有可用股道" in warning["message"] for warning in result["warnings"])


def test_unknown_one_line_can_use_known_selected_line_with_a_warning():
    doc, infra = small_doc()
    infra["members"][0]["UB_TRACKS"] = None
    infra["lines"].append({"LINE_IND": 2})
    infra["members"].append({"node_name": "P", "line_ind": 2, "DB_TRACKS": "[9,10]"})
    selection = TrainChoice(
        train_id="G1",
        lines=[{"line_ind": 1, "directions": ["UB"]}, {"line_ind": 2, "directions": ["DB"]}],
    )
    result = allocate(doc, infra, [selection])
    assert result["can_apply"] and result["changes"][0]["value"] == 9
    assert any("股道资料未知" in warning["message"] for warning in result["warnings"])


@pytest.mark.parametrize("directions", [["UB", "DB"], ["UB", "UB"], []])
def test_same_line_rejects_multiple_or_missing_directions(directions):
    with pytest.raises(ValueError):
        choice("G1", directions)


@pytest.mark.parametrize("kind", ["base", "run"])
def test_partial_unknown_repair_can_be_preflighted_saved_and_exported(tmp_path, kind):
    def unknown_book(result=False):
        book = load_workbook(BytesIO(workbook(result=result)))
        if not result:
            # Q has known DB resources but unknown UB resources.
            book["NODE_INFO"]["I3"] = "[6,9]"
            book["NODE_INFO"]["J3"] = None
        book["TTT1"]["L3"] = 6  # Preserve this fixed occupation when UB is skipped.
        data = BytesIO()
        book.save(data)
        book.close()
        return data.getvalue()

    app = create_app(Settings(database=tmp_path / "db"))
    with TestClient(app) as client:
        ids = branch(client, "partial-unknown")
        if kind == "run":
            create_base(client, ids[2], unknown_book())
            job = run_preview(client, ids[-1], unknown_book(result=True)).json()
        else:
            job = preview(client, ids[2], unknown_book()).json()
        url = f"/api/v1/ingest/jobs/{job['job_id']}/track-repair"
        context = client.get(url + "/context").json()
        selection = [choice(t["id"], (t["direction"],)).model_dump() for t in context["trains"]]
        invalid = deepcopy(selection)
        invalid[0]["lines"][0]["directions"] = ["UB", "DB"]
        rejected = client.post(
            url + "/apply", json={"repair_token": context["repair_token"], "selections": invalid}
        )
        assert rejected.status_code == 422
        repaired = client.post(
            url + "/apply", json={"repair_token": context["repair_token"], "selections": selection}
        )
        assert repaired.status_code == 200 and "spreadsheetml" in repaired.headers["content-type"]
        assert int(repaired.headers["X-Repair-Warnings"]) > 0
        fixed = load_workbook(BytesIO(repaired.content))
        assert fixed["TTT1"]["H3"].value == 9
        assert fixed["TTT1"]["L3"].value == 6
        fixed.close()
        checked = (
            run_preview(client, ids[-1], repaired.content)
            if kind == "run"
            else preview(client, ids[2], repaired.content)
        ).json()
        assert checked["report"]["can_commit"], checked["report"]["issues"]
        saved = (run_commit if kind == "run" else commit)(client, checked)
        assert saved.status_code == 200, saved.text
        object_id = saved.json()["id"]
        object_url = f"/api/v1/{'runs' if kind == 'run' else 'base-timetables'}/{object_id}"
        before = client.get(object_url + "/tables").json()
        repair_url = f"/api/v1/track-repair/{kind}/{object_id}"
        draft = {
            "revision": before["revision"],
            "changes": [
                {
                    "event_id": next(e["id"] for e in before["events"] if e["node_name"] == "P"),
                    "field": "track_code",
                    "value": None,
                }
            ],
        }
        context = client.post(repair_url + "/context", json=draft).json()
        selection = [choice(t["id"], (t["direction"],)).model_dump() for t in context["trains"]]
        response = client.post(
            repair_url + "/apply",
            json=draft | {"repair_token": context["repair_token"], "selections": selection},
        ).json()
        assert response["can_apply"] and response["warnings"]
        assert len(response["changes"]) == 1
        skipped = next(
            e for e in before["events"] if e["node_name"] == "Q" and e["track_code"] == 6
        )
        assert skipped["id"] not in {change["event_id"] for change in response["changes"]}
        assert client.get(object_url + "/tables").json() == before
        changes = {"revision": before["revision"], "changes": response["changes"]}
        validation = client.post(object_url + "/events/validate", json=changes).json()
        assert validation["can_save"], validation["issues"]
        persisted = client.post(
            object_url + "/events/edit",
            json=changes
            | {
                "validation_token": validation["validation_token"],
                "warnings_confirmed": True,
                "train_dates_confirmed": True,
            },
        )
        assert persisted.status_code == 200, persisted.text
        exported = client.post(
            object_url + "/export",
            params={"revision": persisted.json()["revision"], "sheets": "TTT"},
        )
        out = load_workbook(BytesIO(exported.content))
        sheet = out.active
        headers = [cell.value for cell in sheet[1]]
        row = next(
            values
            for values in sheet.iter_rows(min_row=2, values_only=True)
            if values[headers.index("NODE_NAME")] == "Q"
        )
        assert row[headers.index("G2/G3_G")] == 6 and row[headers.index("G1_G")] == 9
        out.close()


def test_missing_g_preflight_reports_each_train_and_correct_header_without_sequence_noise():
    source = workbook(missing=True)
    structure = inspect_workbook(source)
    assert structure["sheets"][0]["trains"] == ["G1", "G2/G3"]
    report = parse_base(source, date(2025, 5, 25), "rebase", track_lists=True)
    problems = [i for i in report["issues"] if i["field"].endswith("_G")]
    assert [(i["field"], i["cell"]) for i in problems] == [("G1_G", "H1"), ("G2/G3_G", "K1")]
    assert not any(i["field"].endswith("_SEQ") for i in report["issues"])


@pytest.mark.parametrize("kind", ["base", "run"])
def test_import_repair_repreflight_editor_save_export_and_stale_rejection(tmp_path, kind):
    source = workbook(missing=True, result=kind == "run")
    app = create_app(Settings(database=tmp_path / "repair.duckdb"))
    with TestClient(app) as client:
        ids = branch(client, "repair")
        if kind == "run":
            create_base(client, ids[2], workbook())
            response = run_preview(client, ids[-1], source)
        else:
            response = preview(client, ids[2], source)
        assert response.status_code == 201, response.text
        job = response.json()
        assert not job["report"]["can_commit"]
        url = f"/api/v1/ingest/jobs/{job['job_id']}/track-repair"
        response = client.get(url + "/context")
        assert response.status_code == 200, response.text
        ctx = response.json()
        assert [t["direction"] for t in ctx["trains"]] == ["DB", "UB"]
        assert all(t["missing_g"] and t["related_lines"] == [1] for t in ctx["trains"])
        body = {
            "repair_token": ctx["repair_token"],
            "selections": [choice(t["id"], (t["direction"],)).model_dump() for t in ctx["trains"]],
        }
        assert client.post(url + "/apply", json=body | {"repair_token": "stale"}).status_code == 409
        partial = client.post(url + "/apply", json=body | {"selections": body["selections"][:1]})
        assert partial.status_code == 200, partial.text
        assert any(
            i.get("field") == "G2/G3_G"
            for i in parse_base(
                partial.content,
                date(2025, 5, 25),
                "rebase",
                track_lists=True,
                result_context={} if kind == "run" else None,
            )["issues"]
        )
        fixed = client.post(url + "/apply", json=body)
        assert fixed.status_code == 200 and "spreadsheetml" in fixed.headers["content-type"], (
            fixed.text
        )
        book = load_workbook(BytesIO(fixed.content))
        table = book["TTT1"]
        assert [c.value for c in table[1]][5:] == [
            f"{c}_{f}" for c in ("G1", "G2/G3") for f in ("A", "D", "G", "SEQ")
        ]
        assert table["H2"].value == 9 and table["L2"].value == 6
        assert table["I2"].value == table["M2"].value == 41
        book.close()
        # Original staged source and failed preflight are untouched.
        assert client.get(url + "/context").json()["repair_token"] == ctx["repair_token"]
        second = (
            run_preview(client, ids[-1], fixed.content).json()
            if kind == "run"
            else preview(client, ids[2], fixed.content).json()
        )
        assert second["report"]["can_commit"], second["report"]["issues"]
        saved = (run_commit if kind == "run" else commit)(client, second)
        assert saved.status_code == 200, saved.text
        object_id = saved.json()["id"]
        object_url = f"/api/v1/{'runs' if kind == 'run' else 'base-timetables'}/{object_id}"
        data = client.get(object_url + "/tables").json()
        before = deepcopy(data["events"])
        edit_url = f"/api/v1/track-repair/{kind}/{object_id}"
        pending = {
            "revision": data["revision"],
            "changes": [
                {"event_id": data["events"][0]["id"], "field": "track_code", "value": None}
            ],
        }
        ctx = client.post(edit_url + "/context", json=pending).json()
        assert "repair_token" in ctx, ctx
        choices = [choice(t["id"], (t["direction"],)).model_dump() for t in ctx["trains"]]
        applied = client.post(
            edit_url + "/apply",
            json=pending | {"repair_token": ctx["repair_token"], "selections": choices},
        )
        assert applied.status_code == 200 and applied.json()["can_apply"], applied.text
        assert client.get(object_url + "/tables").json()["events"] == before
        changes = applied.json()["changes"]
        assert len(changes) == 1
        edit = {"revision": data["revision"], "changes": changes}
        validation = client.post(object_url + "/events/validate", json=edit).json()
        assert validation["can_save"], validation
        saved = client.post(
            object_url + "/events/edit",
            json=edit
            | {
                "validation_token": validation["validation_token"],
                "warnings_confirmed": True,
                "train_dates_confirmed": True,
            },
        )
        assert saved.status_code == 200, saved.text
        assert (
            client.post(
                edit_url + "/apply",
                json=pending | {"repair_token": ctx["repair_token"], "selections": choices},
            ).status_code
            == 409
        )
        latest = saved.json()
        exported = client.post(
            object_url + "/export", params={"revision": latest["revision"], "sheets": "TTT"}
        )
        assert exported.status_code == 200, exported.text
        out = load_workbook(BytesIO(exported.content))
        assert any(cell.value == "G1_G" for cell in out.active[1])
        out.close()


def test_default_only_fills_missing_and_reports_existing_conflicts_without_reassigning():
    doc, infra = small_doc()
    doc["events"][0]["track_code"] = 6
    result = allocate(doc, infra, [choice("G1"), choice("G3")])
    assert result["changes"] == [{"event_id": "G3", "field": "track_code", "value": 7}]
    doc["events"][1]["track_code"] = 6
    before = deepcopy(doc)
    result = allocate(doc, infra, [choice("G1"), choice("G3")])
    assert result["can_apply"] and result["changes"] == []
    assert any(
        "G3" in w["message"] and "G1" in w["message"] and "股道6" in w["message"]
        for w in result["warnings"]
    )
    assert doc == before


def test_conflict_opt_in_keeps_first_window_train_even_if_events_are_in_reverse_order():
    doc, infra = small_doc()
    for event in doc["events"]:
        event["track_code"] = 6
    doc["events"].reverse()
    before = deepcopy(doc)
    result = allocate(doc, infra, [choice("G3"), choice("G1")], repair_conflicts=True)
    assert result["changes"] == [{"event_id": "G3", "field": "track_code", "value": 7}]
    assert doc == before
    infra["lines"][0]["LINE_NAME_CHN"] = "线路"
    context = repair_context(doc, infra)
    assert [t["id"] for t in context["trains"]] == ["G1", "G3"]
    assert all(t["conflict_count"] == 1 for t in context["trains"])
    assert context["warnings"] and doc == before


def test_conflict_repair_does_not_move_a_later_visit_that_becomes_nonconflicting():
    doc, infra = small_doc()
    doc["events"][0].update(track_code=6, arrival={"seconds": 0}, departure={"seconds": 10})
    doc["events"][1].update(track_code=6, arrival={"seconds": 5}, departure={"seconds": 15})
    doc["trains"].append({"id": "G5", "train_course_id": "G5"})
    doc["events"].append(
        doc["events"][1]
        | {
            "id": "G5",
            "train_id": "G5",
            "train_course_id": "G5",
            "arrival": {"seconds": 12},
            "departure": {"seconds": 20},
        }
    )
    result = allocate(doc, infra, [choice("G1"), choice("G3"), choice("G5")], True)
    assert result["changes"] == [{"event_id": "G3", "field": "track_code", "value": 7}]


def test_conflict_repair_protects_later_existing_tracks_and_unselected_visits():
    doc, infra = small_doc()
    for event in doc["events"]:
        event["track_code"] = 6
    doc["trains"].append({"id": "G5", "train_course_id": "G5"})
    doc["events"].append(
        doc["events"][1]
        | {
            "id": "G5",
            "train_id": "G5",
            "train_course_id": "G5",
            "track_code": 7,
        }
    )
    result = allocate(doc, infra, [choice("G1"), choice("G3"), choice("G5")], True)
    assert result["changes"] == [{"event_id": "G3", "field": "track_code", "value": 8}]
    result = allocate(doc, infra, [choice("G1")], True)
    assert result["changes"] == []
    assert any("股道占用冲突" in w["message"] for w in result["warnings"])


def test_capacity_warning_preserves_failed_conflict_and_allows_other_station_repairs():
    doc, infra = small_doc()
    for event in doc["events"]:
        event["track_code"] = 6
    infra["members"][0]["UB_TRACKS"] = "[6]"
    doc["nodes"].append({"node_name": "Q", "node_name_chn": "乙"})
    infra["members"].append({"node_name": "Q", "line_ind": 1, "UB_TRACKS": "[9]"})
    doc["events"].append(doc["events"][0] | {"id": "missing", "node_name": "Q", "track_code": None})
    result = allocate(doc, infra, [choice("G1"), choice("G3")], True)
    assert result["can_apply"] and not result["issues"]
    assert result["changes"] == [{"event_id": "missing", "field": "track_code", "value": 9}]
    assert any("容量不足" in w["message"] for w in result["warnings"])


def test_nonempty_invalid_track_is_warned_and_preserved_instead_of_treated_as_missing():
    doc, infra = small_doc()
    doc["events"][0]["track_code"] = "invalid"
    result = allocate(doc, infra, [choice("G1"), choice("G3")])
    assert result["can_apply"]
    assert {c["event_id"] for c in result["changes"]} == {"G3"}
    assert any("保留已有股道" in w["message"] for w in result["warnings"])


@pytest.mark.parametrize("kind", ["base", "run"])
def test_existing_conflict_mode_roundtrips_through_import_and_editor_without_touching_winner(
    tmp_path, kind
):
    def occupied_book(result=False):
        book = load_workbook(BytesIO(workbook(result=result)))
        for cell, value in (("H2", 9), ("L2", 6), ("H3", 6), ("L3", 6)):
            book["TTT1"][cell] = value
        data = BytesIO()
        book.save(data)
        book.close()
        return data.getvalue()

    app = create_app(Settings(database=tmp_path / "db"))
    with TestClient(app) as client:
        ids = branch(client, "existing-conflicts")
        if kind == "run":
            create_base(client, ids[2], occupied_book())
            job = run_preview(client, ids[-1], occupied_book(result=True)).json()
        else:
            job = preview(client, ids[2], occupied_book()).json()
        repair_url = f"/api/v1/ingest/jobs/{job['job_id']}/track-repair"
        context = client.get(repair_url + "/context").json()
        assert context["warnings"] and all(t["conflict_count"] for t in context["trains"])
        selections = [choice(t["id"], (t["direction"],)).model_dump() for t in context["trains"]]
        body = {"repair_token": context["repair_token"], "selections": selections}
        for enabled in (False, True):
            response = client.post(repair_url + "/apply", json=body | {"repair_conflicts": enabled})
            assert response.status_code == 200, response.text
            book = load_workbook(BytesIO(response.content))
            assert book["TTT1"]["H3"].value == 6
            assert book["TTT1"]["L3"].value == (7 if enabled else 6)
            book.close()
        # Applying either mode only returns a copy; the pending original is unchanged.
        saved = (run_commit if kind == "run" else commit)(client, job)
        assert saved.status_code == 200, saved.text
        object_url = (
            f"/api/v1/{'runs' if kind == 'run' else 'base-timetables'}/{saved.json()['id']}"
        )
        before = client.get(object_url + "/tables").json()
        draft = {"revision": before["revision"], "changes": []}
        url = f"/api/v1/track-repair/{kind}/{saved.json()['id']}"
        context = client.post(url + "/context", json=draft).json()
        selections = [choice(t["id"], (t["direction"],)).model_dump() for t in context["trains"]]
        request = draft | {"repair_token": context["repair_token"], "selections": selections}
        normal = client.post(url + "/apply", json=request).json()
        assert normal["can_apply"] and normal["changes"] == [] and normal["warnings"]
        fixed = client.post(url + "/apply", json=request | {"repair_conflicts": True}).json()
        assert fixed["can_apply"] and len(fixed["changes"]) == 1
        winner, loser = [t["id"] for t in context["trains"]]
        event_by_id = {event["id"]: event for event in before["events"]}
        assert event_by_id[fixed["changes"][0]["event_id"]]["train_id"] == loser
        assert client.get(object_url + "/tables").json() == before
        edit = {"revision": before["revision"], "changes": fixed["changes"]}
        checked = client.post(object_url + "/events/validate", json=edit).json()
        assert checked["can_save"], checked["issues"]
        persisted = client.post(
            object_url + "/events/edit",
            json=edit
            | {
                "validation_token": checked["validation_token"],
                "warnings_confirmed": True,
                "train_dates_confirmed": True,
            },
        )
        assert persisted.status_code == 200, persisted.text
        latest = client.get(object_url + "/tables").json()
        winning_event = next(
            e for e in latest["events"] if e["train_id"] == winner and e["node_name"] == "Q"
        )
        assert winning_event["track_code"] == 6
        changed = next(e for e in latest["events"] if e["id"] == fixed["changes"][0]["event_id"])
        assert changed["track_code"] == fixed["changes"][0]["value"]


@pytest.mark.parametrize("count", [500, 501])
def test_nonoverlapping_missing_visits_do_not_consume_backtracking_budget(count):
    doc, infra = small_doc()
    first = doc["events"][0]
    doc["events"] = [
        first
        | {
            "id": f"visit-{i}",
            "visit_seq": i + 1,
            "arrival": {"seconds": i * 3},
            "departure": {"seconds": i * 3 + 2},
        }
        for i in range(count)
    ]
    result = allocate(doc, infra, [choice("G1")])
    assert result["can_apply"] and len(result["changes"]) == 500
    assert len(result["warnings"]) == count - 500
    assert all(change["value"] == 6 for change in result["changes"])
